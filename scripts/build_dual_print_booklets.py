"""A5/A6/A7 dual-light manuals with legacy and expanded parallax editions."""
import argparse
from collections import Counter
from copy import deepcopy
import hashlib
import importlib.util
import json
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree
from PIL import Image, ImageDraw
import pdfplumber

import build_print_booklets as base
import update_booklet_support as support
from dual_booklet_content import book, parallax_book, illustrated_book
from manual_layout import el, font, inline, picture
from render_manuals import convert_with_word
from verify_print_booklets import normalized

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / '.work/dual-print-booklets'
OUTPUT = ROOT / 'artifacts/print-booklets/dual'
TOTAL = 28
EDITION = 'base'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'v': 'urn:schemas-microsoft-com:vml'}
CONFIG = {
    'a5': dict(w=148.5,h=210,pw=297,ph=210,rows=1,cols=1,out=14,inside=15,text=119.5,
               body_y=22,body_h=169,header=9,footer=196,font=11.5,line=17,title=18,sub=12,small=8.5,
               gap=7,fig=85),
    'a6': dict(w=105,h=148.5,pw=210,ph=297,rows=2,cols=1,out=8,inside=9,text=88,
               body_y=14,body_h=123,header=5.5,footer=140,font=9.5,line=13.2,title=14,sub=10.5,small=7,
               gap=4,fig=57),
    'a7': dict(w=74.25,h=105,pw=297,ph=210,rows=2,cols=2,out=6,inside=7,text=61.25,
               body_y=10,body_h=84,header=4,footer=98,font=7.5,line=9.5,title=11.5,sub=8.5,small=6,
               gap=2,fig=40),
}


def stem(size):
    copies = {'a5':'单套','a6':'双套','a7':'四套'}[size]
    if EDITION == 'illustrated':
        return f'BananaThermal_3.2寸双光_随附说明书_{size.upper()}成册_A4{copies}拼版_v1.8_36面_视差图解版'
    if EDITION == 'parallax':
        return f'BananaThermal_3.2寸双光_随附说明书_{size.upper()}成册_A4{copies}拼版_v1.8_32面_视差详解与测距选配版'
    return f'BananaThermal_3.2寸双光_随附说明书_{size.upper()}成册_A4{copies}拼版_v1.8_28面_仅封面装订虚线版'


def select_edition(edition):
    global EDITION, TOTAL, WORK
    if edition not in ('base','parallax','illustrated'):
        raise ValueError(edition)
    EDITION = edition
    TOTAL = {'base':28,'parallax':32,'illustrated':36}[edition]
    WORK = ROOT / {'base':'.work/dual-print-booklets','parallax':'.work/dual-parallax-booklets',
                   'illustrated':'.work/dual-illustrated-booklets'}[edition]


def draft(size):
    return WORK / size / (stem(size) + '.docx')


def crop_hardware_art():
    """Crop the supplied line art; irregular edges exclude overlapping neighbours.

    Every retained pixel is from the source drawing; no part or label is redrawn.
    """
    source = ROOT / 'docs/assets/user-guide/device-controls.png'
    regions = {
        'controls': [(0,130),(660,130),(660,525),(320,755),(0,755)],
        'rear': [(100,810),(485,675),(540,675),(780,820),(780,1010),(900,1010),(900,1115),(100,1115)],
    }
    with Image.open(source).convert('RGB') as original:
        assert original.size == (1995,1257)
        for name, polygon in regions.items():
            mask = Image.new('L',original.size,0)
            ImageDraw.Draw(mask).polygon(polygon,fill=255)
            result = Image.new('RGB',original.size,'white')
            result.paste(original,(0,0),mask)
            box = mask.getbbox()
            result.crop(box).save(ROOT/f'docs/assets/devices/dual-32-{name}-manual.png')


def para(doc, text, cfg, *, size=None, bold=False, after=None, before=0, style=None, center=False):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(cfg['gap'] if after is None else after)
    pf.line_spacing = Pt(cfg['line'] if size is None else max(size + 2, cfg['line'] * .72))
    pf.widow_control = False
    pf.keep_with_next = False
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inline(p, text, size or cfg['font'])
    if bold:
        for run in p.runs:
            run.bold = True
    return p


def take(doc, start):
    nodes = list(doc._element.body)[start:-1]
    for node in nodes:
        doc._element.body.remove(node)
    return nodes


def table(doc, headers, rows, widths, cfg):
    base.TEXT_W = cfg['text']
    start = len(doc._element.body) - 1
    base.reference_table(doc, headers, rows, widths)
    nodes = list(doc._element.body)[start:-1]
    for node in nodes:
        for spacing in node.iter(qn('w:spacing')):
            spacing.set(qn('w:line'), str(round(cfg['line'] * .98 * 20)))
        for sz in node.iter(qn('w:sz')):
            sz.set(qn('w:val'), str(round((cfg['font'] - .5) * 2)))
        for margin in node.iter(qn('w:tblCellMar')):
            for child in margin:
                child.set(qn('w:w'), '25' if cfg['w'] < 80 else '55')
    return nodes


def figure(doc, name, caption, crop, cfg):
    width = cfg['fig']
    if 'parallax-' in name:
        width = cfg['text'] * .98
    elif 'thermal-view' in name:
        width = cfg['text'] * .86
    elif 'alignment' in name or 'hysteresis' in name:
        width = cfg['text'] * .98
    elif crop and crop[0] > .6:  # USB view
        width *= .78 if cfg['w'] < 80 else .88
    elif 'rear-manual' in name or (crop and crop[1] > .5):  # rear view
        width *= 1.23
    picture(doc, name, caption, width / 25.4, crop)
    for p in list(doc._element.body)[-3:-1]:
        if p.tag == qn('w:p'):
            props = p.find(qn('w:pPr'))
            if props is not None:
                for spacing in props.findall(qn('w:spacing')):
                    spacing.set(qn('w:before'), '0')
                    spacing.set(qn('w:after'), str(round(cfg['gap'] * 20)))
            if ''.join(p.itertext()) and not list(p.iter(qn('w:drawing'))):
                for sz in p.iter(qn('w:sz')):
                    sz.set(qn('w:val'), str(round(cfg['small'] * 2)))


def body_nodes(doc, item, cfg):
    start = len(doc._element.body) - 1
    para(doc, item['chapter'], cfg, size=cfg['small'], bold=True, after=cfg['gap'])
    para(doc, item['title'], cfg, size=cfg['title'], bold=True, style='Heading 1')
    para(doc, item['lead'], cfg)
    step = 0
    for kind, *values in item['blocks']:
        if kind == 'p':
            para(doc, values[0], cfg)
        elif kind == 'h':
            para(doc, values[0], cfg, size=cfg['sub'], bold=True, before=cfg['gap'] / 2)
        elif kind == 's':
            step += 1
            p = para(doc, f'{step:02d}  **{values[0]}**', cfg, after=cfg['gap'] / 2)
            p.paragraph_format.keep_with_next = True
            para(doc, values[1], cfg)
        elif kind == 'i':
            name, caption, _, crop = values
            figure(doc, name, caption, crop, cfg)
        elif kind == 't':
            table(doc, *values, cfg)
        elif kind == 'toc':
            for label, pages in values[0]:
                p = para(doc, label + '\t' + pages, cfg)
                p.paragraph_format.tab_stops.add_tab_stop(Mm(cfg['text'] - .5), WD_TAB_ALIGNMENT.RIGHT)
        else:
            raise ValueError(kind)
    return take(doc, start)


def furniture(doc, cfg, number=None):
    start = len(doc._element.body) - 1
    p = para(doc, 'BANANA THERMAL  3.2 寸双光' if number is None else f'使用说明书\t{number:02d}',
             cfg, size=cfg['small'], after=0)
    p.paragraph_format.tab_stops.add_tab_stop(Mm(cfg['text'] - .5), WD_TAB_ALIGNMENT.RIGHT)
    return take(doc, start)


def cover(doc, cfg, back=False):
    start = len(doc._element.body) - 1
    p = para(doc, 'BANANA THERMAL', cfg, size=cfg['sub'], bold=True,
             before=cfg['gap'], after=cfg['gap'] * 2)
    for r in p.runs:
        r.font.color.rgb = RGBColor.from_string(base.ORANGE)
    if back:
        para(doc, '3.2 寸双光热成像', cfg, bold=True)
        para(doc, '操作速记', cfg, size=cfg['title'], bold=True, before=cfg['gap'] * 5)
        for text in ('**A**  长按换伪彩，短按整理标记。', '**B**  短按拍照，再按返回；长按相册。',
                     '**拨片**  左右选择，按下确认。', '**快捷键**  打开常用参数菜单。'):
            para(doc, text, cfg)
        para(doc, '详细操作见第 5–8 页。', cfg, size=cfg['small'])
        para(doc, '更多资料与交流支持见封三', cfg, bold=True, before=cfg['gap'] * 5)
        para(doc, 'QQ群 1002979587\n扫码与开源项目地址见内封底。', cfg, size=cfg['small'])
    else:
        para(doc, '双光热成像\n使用说明书', cfg,
             size={'a5':25,'a6':20,'a7':16}[cfg['size']], bold=True, style='Title')
        para(doc, '3.2 寸', cfg, size=cfg['title'], bold=True)
        figure(doc, 'manual-design/cover-dual.png', '外观示意  部件见第 2–4 页', (0,0,.31,0),
               {**cfg, 'fig': cfg['text'] * .93})
        para(doc, '设备操作  双光对齐  测温设置', cfg, bold=True, before=cfg['gap'] * 2)
        para(doc, '请在使用前阅读，并随设备妥善保留。', cfg)
    para(doc, '随附说明书 v1.8  ·  2026 年 10 月', cfg, size=cfg['small'])
    return take(doc, start)


def resources(doc, cfg):
    start = len(doc._element.body) - 1
    para(doc, '资料与交流支持', cfg, size=cfg['title'], bold=True)
    para(doc, '扫码加入 QQ 群，获取更多资料与交流支持。', cfg)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(cfg['gap'])
    p.paragraph_format.line_spacing = 1
    width = {'a5':44,'a6':36,'a7':30}[cfg['size']]
    pic = p.add_run().add_picture(str(support.ASSET), width=Mm(width))
    with Image.open(support.ASSET) as im:
        pic.height = Mm(width * im.height * (1-.277-.239) / (im.width * (1-.07-.07)))
    pic._inline.xpath('.//pic:blipFill')[0].insert(1, el('a:srcRect', l=7000,t=27700,r=7000,b=23900))
    pic._inline.docPr.set('descr', 'QQ群 1002979587 加群二维码')
    para(doc, '**QQ群 1002979587**', cfg, center=True)
    for label, repo in support.PROJECTS:
        para(doc, label, cfg, size=cfg['sub'], bold=True, after=0)
        para(doc, support.PREFIX + '\n' + repo, cfg,
             size=8 if cfg['size']=='a7' else cfg['font'] - .5,
             after=1 if cfg['size']=='a7' else cfg['gap'])
    return take(doc, start)


def build(size):
    crop_hardware_art()
    cfg = {**CONFIG[size], 'size':size}
    plan = {'base':book,'parallax':parallax_book,'illustrated':illustrated_book}[EDITION]()
    assert len(plan) == TOTAL-4
    # This is a new manual built with the established native-page system.
    doc = Document()
    base.setup(doc, '32')
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT if cfg['ph'] > cfg['pw'] else WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(cfg['pw']), Mm(cfg['ph'])
    for name in ('Normal', 'Caption', 'Title', 'Heading 1', 'Heading 2'):
        st = doc.styles[name]
        st.font.size = Pt(cfg['font'])
        st.paragraph_format.line_spacing = Pt(cfg['line'])
        st.paragraph_format.space_before = st.paragraph_format.space_after = Pt(0)
    doc.core_properties.title = 'Banana Thermal 3.2寸双光热成像随附说明书'
    doc.core_properties.subject = f'{size.upper()} 成册 A4 预拼版 {TOTAL} 面 仅封面装订虚线'
    doc.core_properties.keywords = f'dual; {size}; saddle stitch; v1.8'
    contents = {1:cover(doc,cfg), TOTAL:cover(doc,cfg,True), TOTAL-1:resources(doc,cfg)}
    heads, feet = {}, {}
    for number, item in enumerate(plan, 1):
        contents[number+2] = body_nodes(doc,item,cfg)
        heads[number+2] = furniture(doc,cfg)
        feet[number+2] = furniture(doc,cfg,number)
    sid = 30000
    for face, pair in enumerate(base.pairs(TOTAL), 1):
        anchor = doc.add_paragraph()
        pf = anchor.paragraph_format
        pf.page_break_before = face > 1
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = Pt(1)
        for row in range(cfg['rows']):
            for col in range(cfg['cols']):
                ox, oy = col*2*cfg['w'], row*cfg['h']
                if face == 1:
                    support.guide(anchor._p, ox+cfg['w'], oy+5, ox+cfg['w'], oy+cfg['h']-5,
                                  f'FoldCover_{row}_{col}', True)
                for half, logical in enumerate(pair):
                    if logical == 2:
                        continue
                    x = ox+half*cfg['w']+(cfg['out'] if half == 0 else cfg['inside'])
                    if logical in heads:
                        for nodes, y, height in ((heads[logical],cfg['header'],5),
                                                  (feet[logical],cfg['footer'],5)):
                            sid += 1
                            base.box(anchor,deepcopy(nodes),x,oy+y,cfg['text'],height,sid)
                        y, height = cfg['body_y'], cfg['body_h']
                    else:
                        y, height = (16,180) if size=='a5' else ((10,131) if size=='a6' else (7,91))
                    sid += 1
                    base.box(anchor,deepcopy(contents[logical]),x,oy+y,cfg['text'],height,sid)
    target = draft(size)
    target.parent.mkdir(parents=True,exist_ok=True)
    doc.save(target)
    evidence = {'size':size,'edition':EDITION,'logical_pages':TOTAL,'body_pages':len(plan),'a4_faces':TOTAL//2,'a4_sheets':TOTAL//4,
                'copies':cfg['rows']*cfg['cols'],'page_pairs':base.pairs(TOTAL),
                'source_basis':['docs/user-guide/dual-light.md','docs/menu/','../3inch2-Heimann-dual-thermal/src/'],
                'existing_docx_hashes':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest()
                    for p in (ROOT/'artifacts').rglob('*.docx')}}
    (target.parent/'evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
    print(size, 'built', TOTAL, 'finished faces')


def render(size, renderer_path):
    target = draft(size)
    spec = importlib.util.spec_from_file_location('packaged_renderer',renderer_path)
    renderer = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(renderer)
    renderer.convert_to_pdf = convert_with_word
    pages = renderer.rasterize(str(target),str(target.parent/'render'),160,verbose=True,emit_pdf=True)
    print(size,len(pages),'A4 faces')


def verify(size):
    target = draft(size);cfg = CONFIG[size];errors=[]
    check = lambda val,msg: errors.append(msg) if not val else None
    doc = Document(target)
    text = ''.join(doc._element.body.itertext())
    for phrase in ('不要','不能','不是','不支持','没有','不应','无法','不可以','单光','standard'):
        check(phrase not in text,'forbidden copy '+phrase)
    for phrase in ('1002979587','NextBananaThermalStudio','BananaThermal-Studio','BananaThermalFirmware'):
        check(phrase in text,'missing resource '+phrase)
    anchors = doc._element.body.findall(qn('w:p'))
    check(len(anchors)==TOTAL//2,'A4 face anchors')
    for face, anchor in enumerate(anchors,1):
        lines = anchor.findall('.//v:line',NS)
        check(bool(lines)==(face==1),'guide location '+str(face))
        check(all(l.get('id','').startswith('FoldCover') for l in lines),'non-cover guide')
    with pdfplumber.open(target.parent/'render'/(target.stem+'.pdf')) as pdf:
        check(len(pdf.pages)==TOTAL//2,'rendered A4 count')
        for face,(page,pair,anchor) in enumerate(zip(pdf.pages,base.pairs(TOTAL),anchors),1):
            for row in range(cfg['rows']):
                for col in range(cfg['cols']*2):
                    x0=col*cfg['w']/25.4*72;y0=row*cfg['h']/25.4*72
                    bounds=(x0,y0,min(x0+cfg['w']/25.4*72,page.width),min(y0+cfg['h']/25.4*72,page.height))
                    crop=page.crop(bounds);native=[]
                    for shape in anchor.findall('.//v:shape',NS):
                        style=shape.get('style','')
                        x=float(re.search(r'margin-left:([\d.]+)pt',style)[1])
                        y=float(re.search(r'margin-top:([\d.]+)pt',style)[1])
                        if x0 <= x < bounds[2] and y0 <= y < bounds[3]:
                            native.extend(t.text or '' for t in shape.iter(qn('w:t')))
                    expected=Counter(normalized(''.join(native)))
                    actual=Counter(normalized(crop.extract_text(x_tolerance=1,y_tolerance=2) or ''))
                    check(expected==actual,f'face {face} cell {row},{col} text difference {dict(expected-actual)}/{dict(actual-expected)}')
                    chars=[c for c in crop.chars if normalized(c['text'])]
                    if chars:
                        check(min(c['x0'] for c in chars)>=x0+3/25.4*72,f'face {face} cell {row},{col} left safety')
                        check(max(c['x1'] for c in chars)<=bounds[2]-3/25.4*72,f'face {face} cell {row},{col} right safety')
                        check(max(c['bottom'] for c in chars)<=bounds[3]-2.8/25.4*72,f'face {face} cell {row},{col} bottom safety')
                        if 3 <= pair[col%2] <= TOTAL-2:
                            bodychars=[c for c in chars if y0+cfg['body_y']/25.4*72 <= c['top'] < y0+cfg['footer']/25.4*72]
                            if bodychars:
                                check(max(c['bottom'] for c in bodychars) <= y0+(cfg['body_y']+cfg['body_h'])/25.4*72,
                                      f'face {face} cell {row},{col} body frame clips text')
                    if pair[col%2]==2:
                        check(not actual and not crop.images,'inside front must stay blank')
                    if pair[col%2]==TOTAL-1:
                        check(len(crop.images)==1,'resource QR')
                    if EDITION == 'illustrated' and pair[col%2] in (18,20,22,24):
                        check(len(crop.images)==1,f'face {face} cell {row},{col} parallax image')
                        for im in crop.images:
                            check(im['x0'] >= x0+cfg['out']/25.4*72-.5 and
                                  im['x1'] <= bounds[2]-cfg['out']/25.4*72+.5,
                                  f'face {face} cell {row},{col} parallax image width')
                            check(im['bottom'] <= y0+(cfg['body_y']+cfg['body_h'])/25.4*72,
                                  f'face {face} cell {row},{col} parallax image clips')
    ev=json.loads((target.parent/'evidence.json').read_text(encoding='utf-8'))
    for path,digest in ev['existing_docx_hashes'].items():
        check(hashlib.sha256((ROOT/path).read_bytes()).hexdigest()==digest,'existing document changed '+path)
    # Sample cover, hardware, fusion, alarms, and resource side; full text is machine-checked above.
    samples=([1,2,13,15,17,18] if EDITION == 'illustrated' else
             ([1,2,10,11,13,14,15,16] if EDITION == 'parallax' else [1,2,4,6,7,11,14]))
    for face in samples:
        with Image.open(target.parent/'render'/f'page-{face}.png') as im:
            im.crop((0,0,round(im.width/cfg['cols']),round(im.height/cfg['rows']))).save(target.parent/f'sample-{face}.png')
    report={'errors':errors,'sha256':hashlib.sha256(target.read_bytes()).hexdigest(),
            'finished_pages':TOTAL,'a4_faces':TOTAL//2,'sampled_faces':samples,'full_native_text_check':True}
    (target.parent/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(size,json.dumps(report,ensure_ascii=False))
    return errors


def publish(size):
    target=draft(size);report=json.loads((target.parent/'verification.json').read_text(encoding='utf-8'))
    assert not report['errors'], report['errors']
    assert report['sha256']==hashlib.sha256(target.read_bytes()).hexdigest()
    review=json.loads((target.parent/'sample-review.json').read_text(encoding='utf-8'))
    assert review['sha256']==report['sha256'] and review['passed']
    OUTPUT.mkdir(parents=True,exist_ok=True)
    with (OUTPUT/target.name).open('xb') as out:
        out.write(target.read_bytes())
    print(OUTPUT/target.name)


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('action',choices=['build','render','verify','publish'])
    parser.add_argument('--size',choices=['all',*CONFIG],default='all')
    parser.add_argument('--renderer',type=Path)
    parser.add_argument('--edition',choices=['base','parallax','illustrated'],default='base')
    args=parser.parse_args();errors=0
    select_edition(args.edition)
    for size in CONFIG if args.size=='all' else [args.size]:
        if args.action=='render':render(size,args.renderer)
        elif args.action=='verify':errors+=len(verify(size))
        else:globals()[args.action](size)
    raise SystemExit(bool(errors))
