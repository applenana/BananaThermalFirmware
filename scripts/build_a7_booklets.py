#!/usr/bin/env python3
"""Reflow approved A6 manuals into editable A7 books, four copies on A4."""
import argparse
from collections import Counter
from copy import deepcopy
import hashlib
import importlib.util
import json
import math
from pathlib import Path
import re
import shutil
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from lxml import etree
from PIL import Image
import pdfplumber

import build_print_booklets as base
import build_a6_booklets as a6
from print_booklet_content import MODELS, book
from render_manuals import convert_with_word
from verify_print_booklets import normalized

ROOT=base.ROOT
WORK=ROOT/'.work/a7-booklets'
OUTPUT=ROOT/'artifacts/print-booklets/a7'
W,H,TEXT_W=74.25,105,61.25
BODY,LOGICAL=56,60
NS=a6.NS


def stem(model):
    return f"BananaThermal_{MODELS[model]['size']}寸单光_随附说明书_A7成册_A4四套拼版_v1.8"


def draft(model): return WORK/model/(stem(model)+'.docx')


def reference(text):
    def replace(match):
        start=int(match[1]); end=int(match[2] or start)
        return f'第 {2*start-1}–{2*end} 页'
    return re.sub(r'第\s*(\d+)\s*(?:[–—\-至]\s*(\d+)\s*)?页', replace, text)


def remap(value):
    if isinstance(value,str): return reference(value)
    if isinstance(value,tuple): return tuple(remap(v) for v in value)
    if isinstance(value,list): return [remap(v) for v in value]
    if isinstance(value,dict): return {k:remap(v) for k,v in value.items()}
    return value


def text_width(text,size):
    return sum(size*(1 if ord(c)>255 else .51) for c in text.replace('**',''))


def lines(text,size,width=TEXT_W):
    return sum(max(1,math.ceil(text_width(s,size)/(width*72/25.4))) for s in text.split('\n'))


def cost(block):
    kind,*v=block
    if kind=='h': return lines(v[0],10.5)*14+7
    if kind=='p': return lines(v[0],9)*12.5+4
    if kind=='n': return lines(v[0]+'  '+v[1],8.5)*12+7
    if kind=='s': return lines(v[0],9.5)*13+lines(v[1],9,TEXT_W-4)*12.5+8
    if kind=='toc': return 15*len(v[0])
    if kind=='t':
        headers,rows,widths=v
        return sum(max(lines(c,8.5,TEXT_W*w/sum(widths)-2.3) for c,w in zip(row,widths))*11.5+4
                   for row in [headers,*rows])+3
    if kind=='i':
        name,caption,width,crop=v
        with Image.open(ROOT/'docs/assets'/name) as im:
            l,t,r,b=crop or (0,0,0,0)
            ratio=im.height*(1-t-b)/(im.width*(1-l-r))
        image_w=min(width*25.4,TEXT_W)
        return image_w*ratio*72/25.4+lines(caption,7)*10+3
    raise ValueError(kind)


def expand(blocks):
    result=[]
    for b in blocks:
        if b[0]=='t':
            for row in b[2]: result.append(('t',b[1],[row],b[3]))
        elif b[0]=='toc':
            for title,ref in b[1]:
                nums=[int(x) for x in re.findall(r'\d+',ref)]
                result.append(('toc',[(title,f'{2*nums[0]-1}–{2*nums[-1]}')]))
        else: result.append(b)
    return result


def merge(blocks):
    result=[]
    for b in blocks:
        if result and b[0]=='t' and result[-1][0]=='t' and b[1]==result[-1][1]:
            prev=result[-1]; result[-1]=('t',prev[1],prev[2]+b[2],prev[3])
        elif result and b[0]=='toc' and result[-1][0]=='toc':
            result[-1]=('toc',result[-1][1]+b[1])
        else: result.append(b)
    return result


def plan(model):
    result=[]
    for old,item in enumerate(remap(book(model)),1):
        blocks=expand(item['blocks'])
        if old in (2,3):
            # Give the original annotated hardware drawing most of its own page.
            for index,b in enumerate(blocks):
                if b[0]=='i':
                    name,caption,width,crop=b[1:]
                    with Image.open(ROOT/'docs/assets'/name) as im:
                        l,t,r,bot=crop or (0,0,0,0)
                        ratio=im.height*(1-t-bot)/(im.width*(1-l-r))
                    blocks[index]=('i',name,caption,min(TEXT_W,64/ratio)/25.4,crop)
        options=[]
        for split in range(1,len(blocks)):
            if blocks[split-1][0]=='h': continue
            first,second=merge(blocks[:split]),merge(blocks[split:])
            heights=[sum(map(cost,first))+28+lines(item['title'],12.5)*16+lines(item['lead'],9)*12.5,
                     sum(map(cost,second))+28+lines(item['title']+' 续',12.5)*16]
            options.append((max(heights)+abs(heights[0]-heights[1])*.08,first,second,heights))
        _,first,second,heights=min(options,key=lambda x:x[0])
        if old==1:
            split=next(i for i,b in enumerate(blocks) if b[0]=='h')
            first,second=merge(blocks[:split]),merge(blocks[split:])
        elif old in (2,3):
            split=next(i for i,b in enumerate(blocks) if b[0]=='i')+1
            first,second=merge(blocks[:split]),merge(blocks[split:])
        elif old==7:
            split=next(i for i,b in enumerate(blocks) if b[0]=='s')+1
            first,second=merge(blocks[:split]),merge(blocks[split:])
        elif old==28:
            first,second=merge(blocks[:4]),merge(blocks[4:])
        for continuation,blocks in enumerate((first,second)):
            result.append(dict(title=item['title']+(' 续' if continuation else ''),
                               lead=(item['lead'] if continuation else '') if old in (2,3) else ('' if continuation else item['lead']),
                               chapter=item['chapter'],step_offset=sum(b[0]=='s' for b in first) if continuation else 0,
                               blocks=blocks,old=old,continuation=bool(continuation),estimate=round(heights[continuation],1)))
    assert len(result)==BODY
    return result


def restyle(nodes, cover=False):
    sizes={11:9,10.5:8.5,11.5:9.5,12:10.5,13:10.5,14:11,
           18:12.5,19:13,25:16,9:7,8.5:6.5,8:6,10:8}
    for node in nodes:
        for element in node.iter():
            if element.tag in (qn('w:sz'),qn('w:szCs')):
                old=int(element.get(qn('w:val')))/2
                element.set(qn('w:val'),str(round(sizes.get(old,old)*2)))
            elif element.tag==qn('w:spacing'):
                for attr in ('before','after'):
                    key=qn('w:'+attr)
                    if key in element.attrib:
                        element.set(key,str(round(int(element.get(key))*(.15 if cover else .42))))
                if element.get(qn('w:lineRule'))=='exact':
                    element.set(qn('w:line'),str(round(int(element.get(qn('w:line')))*.78)))
            elif element.tag==qn('w:ind'):
                for attr in ('left','right'):
                    key=qn('w:'+attr)
                    if key in element.attrib: element.set(key,str(round(int(element.get(key))*.55)))
            elif element.getparent() is not None and element.getparent().tag==qn('w:tblCellMar'):
                if element.tag in (qn('w:top'),qn('w:bottom')): element.set(qn('w:w'),'40')
                elif element.tag in (qn('w:left'),qn('w:right')): element.set(qn('w:w'),'55')
        # Cover art has a separate size budget; all other illustrations use full text width.
        if cover:
            for element in node.iter():
                if element.tag in (qn('wp:extent'),qn('a:ext')):
                    for attr in ('cx','cy'):
                        if attr in element.attrib: element.set(attr,str(round(int(element.get(attr))*.57)))
        if cover:
            for t in node.iter(qn('w:t')):
                if t.text: t.text=reference(t.text)
    return nodes


def compose(doc,item,number):
    adapted=deepcopy(item)
    blocks=[]
    for block in adapted['blocks']:
        if block[0]=='i' and 'menu-flow' in block[1]:
            blocks.extend([('p',block[2]),('t',['次序','按键'],
                [['1 进入','A+B'],['2 选择','A/B'],['3 确认','A+B'],['4 调值','A/B'],['5 确认','A+B']],[2,2.5])])
        elif block[0]=='i' and 'alarm-hysteresis' in block[1]:
            blocks.extend([('p',block[2]),('t',['温度变化','状态'],
                [['达到 60 ℃','满足延时后触发'],['回落至 58–60 ℃','保持报警'],['降至 58 ℃或以下','满足延时后恢复']],[2.3,2.2])])
        else:blocks.append(block)
    adapted['blocks']=blocks
    nodes=base.compose_body(doc,adapted,number)
    nodes=restyle(nodes)
    for node in nodes:
        for t in node.iter(qn('w:t')):
            if re.fullmatch(r'\d{2}  ',t.text or ''):
                t.text=f'{int(t.text)+item["step_offset"]:02d}  '
    for p in nodes:
        if p.tag==qn('w:p') and not ''.join(p.xpath('.//w:t/text()')) and not p.xpath('.//w:drawing'):
            if p.getparent() is not None: p.getparent().remove(p)
    return [p for p in nodes if p.getparent() is not None]


def guide(anchor,x,y,width,height,ident):
    pict=etree.SubElement(etree.SubElement(anchor,qn('w:r')),qn('w:pict'))
    v=NS['v']; line=etree.SubElement(pict,'{'+v+'}line',nsmap={'v':v})
    line.set('id',ident); line.set('from','0pt,0pt')
    line.set('to',f'{width*72/25.4:.4f}pt,{height*72/25.4:.4f}pt')
    line.set('style',f'position:absolute;margin-left:{x*72/25.4:.4f}pt;margin-top:{y*72/25.4:.4f}pt;'
             f'width:{width*72/25.4:.4f}pt;height:{height*72/25.4:.4f}pt;z-index:9000;'
             'mso-position-horizontal-relative:page;mso-position-vertical-relative:page')
    line.set('strokecolor','#8C8C8C');line.set('strokeweight','0.5pt')
    etree.SubElement(line,'{'+v+'}stroke').set('dashstyle','dash')
    etree.SubElement(line,'{urn:schemas-microsoft-com:office:word}wrap').set('type','none')


def build(model):
    source=a6.OUTPUT/(a6.stem(model)+'.docx')
    doc=Document(source)
    for node in list(doc._element.body):
        if node.tag!=qn('w:sectPr'):doc._element.body.remove(node)
    base.TEXT_W=TEXT_W
    sec=doc.sections[0];sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width=Mm(297);sec.page_height=Mm(210)
    normal=doc.styles['Normal'];normal.font.size=Pt(9);normal.paragraph_format.line_spacing=Pt(12.5)
    for name,size,line in [('Title',16,20),('Heading 1',12.5,16),('Heading 2',10.5,14),('Caption',7,10)]:
        st=doc.styles[name];st.font.size=Pt(size);st.paragraph_format.line_spacing=Pt(line)
        st.paragraph_format.space_before=Pt(3 if name=='Heading 2' else 0);st.paragraph_format.space_after=Pt(3)
    doc.core_properties.subject='A7 成册 A4 横向四套预拼版 60 面 骑马订'
    doc.core_properties.keywords='A7; A4 landscape; four copies; saddle stitch; v1.8'
    pages=plan(model);sid=3000
    # Compose each logical page once, then clone native XML for four copies.
    contents={};headers={};footers={}
    for logical in range(1,LOGICAL+1):
        if logical in (2,LOGICAL-1):continue
        if logical==1: nodes=restyle(base.compose_cover(doc,model),True)
        elif logical==LOGICAL:nodes=restyle(base.compose_back(doc,model),True)
        else:
            number=logical-2;item=pages[number-1];nodes=compose(doc,item,number)
            headers[logical]=restyle(base.furniture(doc,model,item,number))
            footers[logical]=restyle(base.furniture(doc,model,item,number,True))
        contents[logical]=nodes
    for nodes in [*contents.values(),*headers.values(),*footers.values()]:
        for node in nodes:
            if node.getparent() is not None:node.getparent().remove(node)
    for side,pair in enumerate(base.pairs(LOGICAL),1):
        anchor=doc.add_paragraph();pf=anchor.paragraph_format
        pf.page_break_before=side>1;pf.space_before=pf.space_after=Pt(0);pf.line_spacing=Pt(1)
        guide(anchor._p,5,105,287,0,f'A7CutH{side}')
        guide(anchor._p,148.5,5,0,200,f'A7CutV{side}')
        for row in range(2):
            for col in range(2):
                for half,logical in enumerate(pair):
                    if logical in (2,LOGICAL-1):continue
                    x=col*2*W+half*W+(6 if half==0 else 7);oy=row*H
                    if logical not in (1,LOGICAL):
                        for nodes,y,h in [(headers[logical],4,4),(footers[logical],99,4)]:
                            sid+=1;base.box(anchor,deepcopy(nodes),x,oy+y,TEXT_W,h,sid)
                    sid+=1;base.box(anchor,deepcopy(contents[logical]),x,oy+(7 if logical in (1,LOGICAL) else 12),TEXT_W,91 if logical in (1,LOGICAL) else 84,sid)
    target=draft(model);target.parent.mkdir(parents=True,exist_ok=True);doc.save(target)
    evidence=dict(source=str(source),source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                  body_pages=BODY,logical_pages=LOGICAL,a4_sides=30,a4_sheets=15,copies=4,
                  cut_x_mm=148.5,cut_y_mm=105,fold_x_mm=[74.25,222.75],plan=pages,
                  pairs=base.pairs(LOGICAL),before_docx={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest()
                  for p in (ROOT/'artifacts').rglob('*.docx') if OUTPUT not in p.parents and not p.name.startswith('~$')})
    (target.parent/'evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
    print(target)


def render(model,renderer_path):
    spec=importlib.util.spec_from_file_location('packaged_renderer',renderer_path)
    renderer=importlib.util.module_from_spec(spec);spec.loader.exec_module(renderer)
    renderer.convert_to_pdf=convert_with_word
    pages=renderer.rasterize(str(draft(model)),str(WORK/model/'render'),160,verbose=True,emit_pdf=True)
    print(f'{model}: {len(pages)} physical faces')


def verify(model):
    errors=[]
    def check(test,msg):
        if not test:errors.append(msg)
    with ZipFile(draft(model)) as z:root=etree.fromstring(z.read('word/document.xml'))
    anchors=root.xpath('./w:body/w:p',namespaces=NS)
    check(len(anchors)==30,'physical anchors')
    check(len(root.xpath('.//v:line',namespaces=NS))==60,'two cutting guides per face')
    settings=Document(draft(model)).settings.element
    for tag in ('bookFoldPrinting','bookFoldRevPrinting','mirrorMargins','printTwoOnOne'):
        check(not settings.xpath('./w:'+tag),'second imposition '+tag)
    all_text=''.join(root.xpath('.//w:t/text()',namespaces=NS))
    for phrase in ('默认款不带触屏','选配触屏','相关项目需配合触屏硬件使用'):
        check(phrase in all_text,'touch option missing '+phrase)
    for phrase in ('本机没有','本设备没有','没有可见光','快捷菜单','复位孔','联系售后','不要','不能','不是','不代表','不等于'):
        check(phrase not in all_text,'editorial regression '+phrase)
    mapping=base.pairs(LOGICAL)
    for number,item in enumerate(plan(model),1):
        logical=number+2
        face=next(i for i,pair in enumerate(mapping) if logical in pair)
        half=mapping[face].index(logical)
        # First signature only; all eight rendered cells are checked below.
        body=[]
        for shape in anchors[face].xpath('.//v:shape',namespaces=NS):
            style=shape.get('style')
            x=float(re.search(r'margin-left:([\d.]+)pt',style)[1])/(72/25.4)
            y=float(re.search(r'margin-top:([\d.]+)pt',style)[1])/(72/25.4)
            if half*W<x<(half+1)*W and abs(y-12)<.01:body.append(shape)
        check(len(body)==1,f'body frame {number}')
        if body:
            actual=normalized(''.join(body[0].xpath('.//w:t/text()',namespaces=NS)))
            for text in [item['title'],item['lead']]:
                check(normalized(text.replace('**','').replace('{','').replace('}','')) in actual,f'body {number} source lead/title')
    extent=[]
    with pdfplumber.open(WORK/model/'render'/(stem(model)+'.pdf')) as pdf:
        check(len(pdf.pages)==30,'rendered page count')
        for face,(page,pair,anchor) in enumerate(zip(pdf.pages,base.pairs(LOGICAL),anchors),1):
            check(abs(page.width-297/25.4*72)<.2 and abs(page.height-210/25.4*72)<.2,'A4 landscape size')
            shapes=anchor.xpath('.//v:shape',namespaces=NS)
            for row in range(2):
                for col in range(4):
                    left=col*page.width/4;top=row*page.height/2
                    crop=page.crop((left,top,left+page.width/4,top+page.height/2))
                    logical=pair[col%2];label=f'face {face} cell {row},{col} logical {logical}'
                    native=[]
                    for shape in shapes:
                        style=shape.get('style')
                        x=float(re.search(r'margin-left:([\d.]+)pt',style)[1]);y=float(re.search(r'margin-top:([\d.]+)pt',style)[1])
                        if int(x/(page.width/4))==col and int(y/(page.height/2))==row:
                            native.extend(shape.xpath('.//w:t/text()',namespaces=NS))
                    expected=normalized(''.join(native));actual=normalized(crop.extract_text(x_tolerance=1,y_tolerance=2) or '')
                    missing=Counter(expected)-Counter(actual);extra=Counter(actual)-Counter(expected)
                    check(not missing,label+' missing '+str(dict(missing)))
                    check(not extra,label+' extra '+str(dict(extra)))
                    if logical in (2,LOGICAL-1):
                        check(not actual and not crop.images,label+' inside-cover blank');continue
                    chars=[c for c in crop.chars if normalized(c['text'])]
                    check(all(c['x0']>=left+4/25.4*72 and c['x1']<=left+page.width/4-4/25.4*72 for c in chars),label+' horizontal safety')
                    if logical not in (1,LOGICAL):
                        body=[c for c in chars if top+10/25.4*72<c['top']<top+98/25.4*72]
                        low=max((c['bottom']-top for c in body),default=0)/72*25.4
                        check(low<=96,label+' body overflow '+str(round(low,1)))
                        check(f'{logical-2:02d}' in actual,label+' page number')
                        extent.append(dict(body=logical-2,bottom_mm=round(low,1)))
    evidence=json.loads((WORK/model/'evidence.json').read_text(encoding='utf-8'))
    for path,digest in evidence['before_docx'].items():
        if Path(path).name.startswith('~$'):continue  # Word's transient owner/lock record, not a manual.
        check(hashlib.sha256((ROOT/path).read_bytes()).hexdigest()==digest,'existing file modified '+path)
    (WORK/model/'verification.json').write_text(json.dumps(dict(errors=errors,extent=extent,
        draft_sha256=hashlib.sha256(draft(model).read_bytes()).hexdigest()),ensure_ascii=False,indent=2),encoding='utf-8')
    print(model+(': PASS' if not errors else ':\n'+'\n'.join(errors)))
    return errors


def publish(model):
    report=json.loads((WORK/model/'verification.json').read_text(encoding='utf-8'))
    if report['errors']:raise RuntimeError('QA failed')
    if report['draft_sha256']!=hashlib.sha256(draft(model).read_bytes()).hexdigest():raise RuntimeError('Draft changed after QA')
    OUTPUT.mkdir(parents=True,exist_ok=True)
    with draft(model).open('rb') as source,(OUTPUT/(stem(model)+'.docx')).open('xb') as target:shutil.copyfileobj(source,target)


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('action',choices=['build','render','verify','publish'])
    p.add_argument('--model',choices=['all',*MODELS],default='all');p.add_argument('--renderer',type=Path)
    args=p.parse_args();errors=0
    for model in MODELS if args.model=='all' else [args.model]:
        if args.action=='build':build(model)
        elif args.action=='render':render(model,args.renderer)
        elif args.action=='verify':errors+=len(verify(model))
        else:publish(model)
    raise SystemExit(bool(errors))
