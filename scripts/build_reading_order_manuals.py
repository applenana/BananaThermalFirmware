#!/usr/bin/env python3
"""Convert approved imposed manuals to native, one-page-per-leaf Word originals."""
import argparse
from collections import Counter
from copy import deepcopy
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import shutil
import subprocess
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT,WD_SECTION_START
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Mm,Pt
from docx.text.paragraph import Paragraph
from lxml import etree
import pdfplumber
from PIL import Image,ImageDraw

import build_print_booklets as base
import update_booklet_support as support
import build_compact_a7 as compact
from manual_layout import el,font
from render_manuals import convert_with_word
from verify_print_booklets import normalized

ROOT=base.ROOT;WORK=ROOT/'.work/reading-order'
OUTPUT=ROOT/'artifacts/print-booklets/reading-order'
V='{urn:schemas-microsoft-com:vml}'


def config(size):
    return dict(w=105,h=148.5,logical=32,body=28,top=14,bottom=10.5,left=9,right=8,
                cover_top=10,body_y=14,header=5.5,footer=5.5,font=9.5) if size=='a6' else dict(
                w=74.25,h=105,logical=28,body=24,top=10,bottom=7,left=7,right=6,
                cover_top=7,body_y=10,header=4,footer=3.8,font=7.5)


def source_path(size,model):
    return support.paths('a6',model)[2] if size=='a6' else compact.a7.OUTPUT/(compact.stem(model)+'.docx')


def stem(size,model):
    pages=config(size)['logical']
    return f'BananaThermal_{base.MODELS[model]["size"]}寸单光_随附说明书_{size.upper()}原始单页顺序版_{pages}面_v1.8'


def draft(size,model):return WORK/size/model/(stem(size,model)+'.docx')


def extracted_pages(doc,size):
    cfg=config(size);anchors=[n for n in doc._element.body if n.tag==qn('w:p')]
    pages={2:[]};names={}
    for anchor,pair in zip(anchors,base.pairs(cfg['logical'])):
        for shape in anchor.iter(V+'shape'):
            style=shape.get('style','')
            x=float(re.search(r'margin-left:([\d.]+)pt',style)[1])/(72/25.4)
            y=float(re.search(r'margin-top:([\d.]+)pt',style)[1])/(72/25.4)
            if x>=cfg['w']*2 or y>=cfg['h']:continue
            logical=pair[int(x/cfg['w'])]
            target_y=cfg['cover_top'] if logical in (1,cfg['logical']-1,cfg['logical']) else cfg['body_y']
            if abs(y-target_y)>.01:continue
            content=shape.find('.//'+qn('w:txbxContent'))
            assert logical not in pages,(logical,'duplicate')
            pages[logical]=[deepcopy(n) for n in content]
    assert sorted(pages)==list(range(1,cfg['logical']+1)),sorted(pages)
    return pages


def dimensions(section,cfg,cover=False):
    section.orientation=WD_ORIENT.PORTRAIT
    section.page_width=Mm(cfg['w']);section.page_height=Mm(cfg['h'])
    # Uniform page margins make Word's Apply to whole document safe across
    # the three sections; only the header/footer linkage differs by section.
    section.top_margin=Mm(cfg['top'])
    section.bottom_margin=Mm(cfg['bottom'])
    section.left_margin=Mm(cfg['left']);section.right_margin=Mm(cfg['right']);section.gutter=Mm(0)
    section.header_distance=Mm(cfg['header']);section.footer_distance=Mm(cfg['footer'])
    section.different_first_page_header_footer=False


def bare_paragraph(doc,new_page=False):
    p=doc.add_paragraph();pf=p.paragraph_format
    pf.page_break_before=new_page;pf.keep_with_next=False;pf.keep_together=False
    pf.space_before=pf.space_after=Pt(0);pf.line_spacing=Pt(.05)
    font(p.add_run(''),1)
    return p


def append_page(doc,nodes,new_page=False):
    copied=deepcopy(nodes)
    if not copied:bare_paragraph(doc,new_page);return
    if copied[0].tag==qn('w:p'):
        Paragraph(copied[0],doc._body).paragraph_format.page_break_before=new_page
    elif new_page:bare_paragraph(doc,True)
    for node in copied:
        for bookmark in list(node.iter()):
            if bookmark.tag in (qn('w:bookmarkStart'),qn('w:bookmarkEnd')):bookmark.getparent().remove(bookmark)
        doc._element.body.insert(len(doc._element.body)-1,node)


def unlink(section):
    for hf in (section.header,section.footer):
        hf.is_linked_to_previous=False
        for p in hf.paragraphs:p.clear()


def body_furniture(section,cfg,model):
    unlink(section)
    p=section.header.paragraphs[0]
    p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.line_spacing=Pt(8 if cfg['w']<100 else 10)
    font(p.add_run(f'BANANA THERMAL   {base.MODELS[model]["size"]} 寸单光'),6.5 if cfg['w']<100 else 7)
    p=section.footer.paragraphs[0];pf=p.paragraph_format
    pf.space_before=pf.space_after=Pt(0);pf.line_spacing=Pt(8 if cfg['w']<100 else 10)
    pf.tab_stops.add_tab_stop(Mm(cfg['w']-cfg['left']-cfg['right']-1),WD_TAB_ALIGNMENT.RIGHT)
    size=6.5 if cfg['w']<100 else 7.5
    font(p.add_run('使用说明书\t'),size)
    field=el('w:fldSimple',**{'w:instr':' PAGE \\# "00" '})
    run=el('w:r');run.append(el('w:rPr'));run[0].append(el('w:sz',**{'w:val':int(size*2)}))
    t=el('w:t');t.text='01';run.append(t);field.append(run);p._p.append(field)
    start=el('w:pgNumType',**{'w:start':1});section._sectPr.append(start)


def build(size,model):
    cfg=config(size);source=source_path(size,model);doc=Document(source)
    pages=extracted_pages(doc,size)
    for node in list(doc._element.body):
        if node.tag!=qn('w:sectPr'):doc._element.body.remove(node)
    for tag in ('bookFoldPrinting','bookFoldRevPrinting','mirrorMargins','printTwoOnOne','evenAndOddHeaders'):
        for node in doc.settings.element.xpath('./w:'+tag):node.getparent().remove(node)
    doc.settings.element.append(el('w:updateFields',**{'w:val':'true'}))
    initial=doc.sections[0];dimensions(initial,cfg,True);unlink(initial)
    append_page(doc,pages[1]);append_page(doc,[],True)
    section=doc.add_section(WD_SECTION_START.NEW_PAGE);dimensions(section,cfg)
    body_furniture(section,cfg,model)
    for n in range(1,cfg['body']+1):append_page(doc,pages[n+2],n>1)
    section=doc.add_section(WD_SECTION_START.NEW_PAGE);dimensions(section,cfg,True);unlink(section)
    for node in section._sectPr.findall(qn('w:pgNumType')):section._sectPr.remove(node)
    resources=deepcopy(pages[cfg['logical']-1])
    if size=='a7':
        # Uniform top margins add 3 mm versus the old positioned resource frame.
        # Recover that space from paragraph gaps, keeping QR and type sizes intact.
        for node in resources:
            if node.tag==qn('w:p'):
                pf=Paragraph(node,doc._body).paragraph_format
                if pf.space_after and pf.space_after.pt>1:pf.space_after=Pt(1)
    append_page(doc,resources);append_page(doc,pages[cfg['logical']],True)
    doc.core_properties.subject=f'{size.upper()} 原始单页顺序版 每页一个成品面 {cfg["logical"]}面 未拼版'
    doc.core_properties.keywords=f'{size.upper()};reading order;single pages;editable;not imposed'
    target=draft(size,model);target.parent.mkdir(parents=True,exist_ok=True);doc.save(target)
    evidence={'source':str(source),'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
              'page_count':cfg['logical'],'body_count':cfg['body'],
              'page_text':{str(n):''.join(t.text or '' for p in nodes for t in p.iter(qn('w:t'))) for n,nodes in pages.items()},
              'source_media':media_hashes(source)}
    (target.parent/'evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
    print(target)


def media_hashes(path):
    with ZipFile(path) as z:return {n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}


def render(size,model,renderer_path):
    target=draft(size,model)
    spec=importlib.util.spec_from_file_location('packaged_renderer',renderer_path)
    renderer=importlib.util.module_from_spec(spec);spec.loader.exec_module(renderer)
    renderer.convert_to_pdf=convert_with_word
    pages=renderer.rasterize(str(target),str(target.parent/'render'),160,verbose=True,emit_pdf=True)
    print(size,model,len(pages),'reading-order pages')


def wordcheck(size,model):
    cfg=config(size);target=draft(size,model)
    result=subprocess.run(['powershell.exe','-NoProfile','-NonInteractive','-ExecutionPolicy','Bypass',
        '-File',str(ROOT/'scripts/check_word_book_fold.ps1'),'-InputDocx',str(target),
        '-OutputJson',str(target.parent/'book-fold-check.json'),
        '-SheetWidthMm',str(2*cfg['w']),'-SheetHeightMm',str(cfg['h'])],capture_output=True,timeout=180)
    if result.returncode:raise RuntimeError((result.stdout+result.stderr).decode('utf-8',errors='replace'))
    print(size,model,result.stdout.decode('utf-8',errors='replace').strip())


def verify(size,model):
    cfg=config(size);target=draft(size,model)
    ev=json.loads((target.parent/'evidence.json').read_text(encoding='utf-8'));errors=[]
    def check(value,msg):
        if not value:errors.append(msg)
    doc=Document(target)
    check(not list(doc._element.body.iter(qn('w:txbxContent'))),'floating textboxes remain')
    check(not list(doc._element.body.iter(V+'line')),'production guides remain')
    check(media_hashes(target)==ev['source_media'],'source media changed')
    with ZipFile(target) as new,ZipFile(ev['source']) as old:
        for part in ('word/styles.xml','word/fontTable.xml','word/theme/theme1.xml','word/numbering.xml'):
            if part in old.namelist():check(new.read(part)==old.read(part),'preserve-only package part changed '+part)
    check(hashlib.sha256(Path(ev['source']).read_bytes()).hexdigest()==ev['source_sha256'],'original modified')
    for section in doc.sections:
        check(abs(section.page_width.mm-cfg['w'])<.05 and abs(section.page_height.mm-cfg['h'])<.05,'single-page geometry')
    for tag in ('bookFoldPrinting','bookFoldRevPrinting','mirrorMargins','printTwoOnOne'):
        check(not doc.settings.element.xpath('./w:'+tag),'imposition setting remains '+tag)
    check(not doc.settings.element.xpath('./w:bookFoldPrintingSheets'),'booklet-size override remains')
    bookfold=json.loads((target.parent/'book-fold-check.json').read_text(encoding='utf-8-sig'))
    check(bookfold['beforePages']==bookfold['afterPages']==cfg['logical'],'Word Book fold changes page count')
    with pdfplumber.open(target.parent/'render'/(target.stem+'.pdf')) as pdf:
        check(len(pdf.pages)==cfg['logical'],f'rendered pages {len(pdf.pages)}/{cfg["logical"]}')
        for number,page in enumerate(pdf.pages,1):
            actual=Counter(normalized(page.extract_text(x_tolerance=1,y_tolerance=2) or ''))
            expected=Counter(normalized(ev['page_text'].get(str(number),'')))
            missing=expected-actual
            check(not missing,f'page {number} source content missing {dict(missing)}')
            if 3<=number<=cfg['body']+2:
                check(f'{number-2:02d}' in normalized(page.extract_text() or ''),f'page {number} footer number')
                chars=[c for c in page.chars if normalized(c['text'])]
                # Word's native CJK hanging punctuation extends its advance box
                # beyond the text margin; the glyph itself still has a safe inset.
                check(all(c['x0']>=3/25.4*72 and c['x1']<=page.width-3/25.4*72 for c in chars),f'page {number} horizontal safety')
            if number==2:check(not actual and not page.images,'inside front no longer blank')
    cols,rows=(2,2) if size=='a6' else (3,2)
    for start in range(1,cfg['logical']+1,cols*rows):
        files=[target.parent/'render'/f'page-{n}.png' for n in range(start,min(start+cols*rows,cfg['logical']+1))]
        with Image.open(files[0]) as sample:w,h=sample.size
        canvas=Image.new('RGB',(cols*w,rows*(h+18)),'#eeeeee')
        draw=ImageDraw.Draw(canvas)
        for index,path in enumerate(files):
            x=index%cols*w;y=index//cols*(h+18)
            draw.text((x+3,y+2),f'{size} {model} page {start+index}',fill='black')
            with Image.open(path) as im:canvas.paste(im,(x,y+18))
        canvas.save(target.parent/f'proof-{start}.png')
    report={'errors':errors,'sha256':hashlib.sha256(target.read_bytes()).hexdigest()}
    (target.parent/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(size,model,report)
    return errors


def publish(size,model):
    target=draft(size,model);report=json.loads((target.parent/'verification.json').read_text(encoding='utf-8'))
    assert not report['errors'],report['errors']
    assert report['sha256']==hashlib.sha256(target.read_bytes()).hexdigest()
    review=json.loads((target.parent/'visual-review.json').read_text(encoding='utf-8'))
    assert review['sha256']==report['sha256'] and review['page_count']==config(size)['logical'],'visual QA missing'
    OUTPUT.mkdir(parents=True,exist_ok=True)
    with target.open('rb') as src,(OUTPUT/target.name).open('xb') as out:shutil.copyfileobj(src,out)
    print(OUTPUT/target.name)


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('action',choices=['build','render','wordcheck','verify','publish'])
    p.add_argument('--size',choices=['all','a6','a7'],default='all');p.add_argument('--model',choices=['all','24','28','32'],default='all')
    p.add_argument('--renderer',type=Path);args=p.parse_args();errors=0
    for size in ['a6','a7'] if args.size=='all' else [args.size]:
        for model in ['24','28','32'] if args.model=='all' else [args.model]:
            if args.action=='render':render(size,model,args.renderer)
            elif args.action=='verify':errors+=len(verify(size,model))
            else:globals()[args.action](size,model)
    raise SystemExit(bool(errors))
