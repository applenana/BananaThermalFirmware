#!/usr/bin/env python3
"""Check authored text, designed pagination, links and editable DOCX structure.

This is not a replacement for opening every rendered page. Run after PDF export.
"""
from pathlib import Path
import re
import unicodedata
from urllib.parse import unquote
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
import pdfplumber
from manual_layout import MANUALS, ROOT
from manual_content import pages


def normalize(text):
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("−", "-").replace("\u200b", "")
    return re.sub(r"[\s*{}]", "", text)


def authored(item):
    yield item["title"]
    yield item["lead"]
    for kind, *values in item["blocks"]:
        if kind == "i": yield values[1]
        elif kind == "t":
            for row in [values[0], *values[1]]:
                yield from row
        elif kind == "l": yield values[0]
        else: yield from values


def verify_markdown_links():
    def slug(value):
        value = re.sub(r"[`*_]", "", value).lower()
        value = "".join(c for c in value if c in "-_ " or not unicodedata.category(c).startswith(("P", "S")))
        return value.strip().replace(" ", "-")
    count = 0
    for file in [ROOT / "README.md", ROOT / "artifacts/README.md", *ROOT.glob("docs/**/*.md")]:
        for label, target in re.findall(r"!?\[([^\]]*)\]\(([^)]+)\)", file.read_text(encoding="utf-8")):
            target = unquote(target)
            if target.startswith(("https:", "http:", "mailto:", "app:")) or "/releases" in target:
                continue
            path, _, anchor = target.partition("#")
            dest = (file.parent / path).resolve() if path else file
            assert dest.is_file(), (file, "Missing local target", target)
            if anchor and dest.suffix == ".md":
                headings = [slug(x) for x in re.findall(r"^#{1,6}\s+(.+)$", dest.read_text(encoding="utf-8"), re.M)]
                assert anchor in headings, (file, "Missing anchor", target)
            count += 1
    print(f"Markdown: {count} local links and anchors; PASS")


def verify(profile):
    stem = MANUALS[profile][1]
    doc = Document(ROOT / "artifacts" / (stem + ".docx"))
    reader = PdfReader(ROOT / "artifacts" / (stem + ".pdf"))
    assert len(reader.pages) == 20, (profile, "Unexpected PDF page count", len(reader.pages))
    assert len(doc.sections) == 20, (profile, "Unexpected Word section count")
    titles = [p for p in doc.paragraphs if p.style.name == "Title"]
    assert len(titles) == 1
    assert not titles[0]._p.xpath(".//w:pBdr")
    assert len(doc.paragraphs) > 150, "Body must remain editable paragraphs"
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Header"):
        style = doc.styles[name]
        assert str(style.font.color.rgb) == "000000", (profile, name)
        assert not style.element.xpath(".//w:pBdr")
    for table in doc.tables:
        assert table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
        borders = table._tbl.xpath("./w:tblPr/w:tblBorders/*")
        assert len(borders) == 6
        assert all(node.get(qn("w:color")) == "D9D9D9" for node in borders)
    for shape in doc.inline_shapes:
        assert shape._inline.docPr.get("descr"), "Missing image description"
    bookmarks = {node.get(qn("w:name")) for node in doc.element.xpath(".//w:bookmarkStart")}
    anchors = doc.element.xpath(".//w:hyperlink[@w:anchor]")
    assert len(anchors) == 18
    assert all(node.get(qn("w:anchor")) in bookmarks for node in anchors)
    for section in list(doc.sections)[1:]:
        commands = [node.text.strip() for node in section.footer._element.xpath(".//w:instrText")]
        assert commands == ["PAGE", "NUMPAGES"], commands
    for rel in doc.part.rels.values():
        if not rel.is_external: continue
        assert str(rel.target_ref).startswith("https://"), rel.target_ref
        marker = "/BananaThermalFirmware/blob/main/"
        if marker in rel.target_ref:
            assert (ROOT / rel.target_ref.split(marker)[1]).exists(), rel.target_ref
    checks = 0
    with pdfplumber.open(ROOT / "artifacts" / (stem + ".pdf")) as pdf:
        for number, item in enumerate(pages(profile)[2:], 3):
            page = pdf.pages[number-1]
            text = normalize(page.extract_text() or "")
            for expected in authored(item):
                assert normalize(expected) in text, (profile, number, "Text missing or spilled", expected)
                checks += 1
            assert normalize(f"{number} / 20") in text, (profile, number, "Page field mismatch")
        for number, page in enumerate(pdf.pages, 1):
            for word in page.extract_words():
                assert word["x0"] >= 45 and word["x1"] <= 570, (profile, number, "Horizontal overflow", word)
                assert word["top"] >= 15 and word["bottom"] <= 778, (profile, number, "Vertical overflow", word)
    # Check that the exported PDF retains all clickable TOC entries and resources.
    annotations = [a.get_object() for page in reader.pages for a in page.get("/Annots", [])]
    links = [a for a in annotations if a.get("/Subtype") == "/Link"]
    assert len(links) >= 25, (profile, "PDF links missing", len(links))
    page_refs = {page.indirect_reference.idnum: number for number, page in enumerate(reader.pages, 1)}
    toc = [a.get_object() for a in reader.pages[1].get("/Annots", [])]
    targets = [page_refs[a["/Dest"][0].idnum] for a in toc if "/Dest" in a]
    assert targets == list(range(3, 21)), (profile, "Incorrect PDF TOC destinations", targets)
    assert reader.outline, "PDF heading bookmarks missing"
    print(f"{profile}: 20 pages, {checks} content checks, {len(links)} PDF links, "
          f"{len(doc.tables)} native tables, {len(doc.inline_shapes)} figures; PASS")


if __name__ == "__main__":
    verify_markdown_links()
    for profile in MANUALS: verify(profile)
