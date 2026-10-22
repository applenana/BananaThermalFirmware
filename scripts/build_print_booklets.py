#!/usr/bin/env python3
"""Compose three offline booklets as editable, already-imposed A4 DOCX files.

The logical pages are placed in independent native Word text boxes. Tables,
paragraphs and figures remain editable. No screenshot of a text page is used.
"""
from __future__ import annotations
import argparse
import json
from pathlib import Path
from lxml import etree
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from manual_layout import font, inline, picture, el, ROOT
from print_booklet_content import MODELS, BODY_PAGES, book

NS_V = "urn:schemas-microsoft-com:vml"
NS_O = "urn:schemas-microsoft-com:office:office"
NS_W10 = "urn:schemas-microsoft-com:office:word"
PAGE_W, PAGE_H, HALF = 297, 210, 148.5
TEXT_W = 119.5
ORANGE, NAVY, GRAY = "D94B17", "183D50", "5B6570"
LOGICAL = BODY_PAGES + 4
OUTPUT = ROOT / "artifacts/print-booklets"


def pairs(total=LOGICAL):
    if total % 4: raise ValueError("Saddle stitch requires a multiple of four pages")
    return [pair for sheet in range(total // 4) for pair in
            ((total-2*sheet, 1+2*sheet), (2+2*sheet, total-1-2*sheet))]


def stem(model):
    return f"BananaThermal_{MODELS[model]['size']}寸单光_随附说明书_A4骑马订拼版_v1.8_触屏选配说明版"


def para(doc, text="", *, size=11, bold=False, before=0, after=7, color="000000", style=None):
    pp = doc.add_paragraph(style=style)
    pp.paragraph_format.space_before = Pt(before)
    pp.paragraph_format.space_after = Pt(after)
    inline(pp, text, size)
    for run in pp.runs:
        if bold: run.bold = True
        if color != "000000": run.font.color.rgb = RGBColor.from_string(color)
    return pp


def reference_table(doc, headers, rows, widths):
    tb = doc.add_table(rows=0, cols=len(headers)); tb.autofit = False
    widths = [TEXT_W * w/sum(widths) for w in widths]
    for col, width in zip(tb.columns, widths): col.width = Mm(width)
    pr = tb._tbl.tblPr
    pr.append(el("w:tblInd", **{"w:w": 0, "w:type": "dxa"}))
    borders = el("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(el("w:"+side, **{"w:val": "single", "w:sz": 4, "w:color": "D9D9D9"}))
    pr.append(borders)
    margins = el("w:tblCellMar")
    for side, value in (("top", 65), ("bottom", 65), ("left", 90), ("right", 90)):
        margins.append(el("w:"+side, **{"w:w": value, "w:type": "dxa"}))
    pr.append(margins)
    for row_no, values in enumerate([headers, *rows]):
        row = tb.add_row(); props = row._tr.get_or_add_trPr(); props.append(el("w:cantSplit"))
        if not row_no: props.append(el("w:tblHeader"))
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Mm(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell._tc.get_or_add_tcPr().append(el("w:shd", **{"w:fill": NAVY if not row_no else ("F2F5F7" if row_no % 2 == 0 else "FFFFFF")}))
            pp = cell.paragraphs[0]
            pf = pp.paragraph_format
            pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = Pt(14.5)
            pf.keep_with_next = row_no == 0
            if headers[j] == "原生阵列": pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            font(pp.add_run(value.replace("_", "_\u200b")), 10.5, row_no == 0, "FFFFFF" if not row_no else "000000")
    pp = para(doc, "", after=2)
    pp.paragraph_format.line_spacing = Pt(3)


def setup(doc, model):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Mm(PAGE_W), Mm(PAGE_H)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(0)
    section.header_distance = section.footer_distance = Mm(0)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption", "Header", "Footer"):
        st = doc.styles[name]; st.font.name = "Calibri"; st.font.color.rgb = RGBColor(0,0,0)
        st._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        for node in st._element.xpath(".//w:color"):
            for key in ("w:themeColor", "w:themeTint", "w:themeShade"): node.attrib.pop(qn(key), None)
        for node in st._element.xpath(".//w:pBdr"): node.getparent().remove(node)
        st.paragraph_format.widow_control = True
    st = doc.styles["Normal"]; st.font.size = Pt(11)
    st.paragraph_format.line_spacing = Pt(16.5); st.paragraph_format.space_after = Pt(7)
    for name, size, spacing in (("Title", 25, 34), ("Heading 1",18,25), ("Heading 2",12,19)):
        st = doc.styles[name]; st.font.size = Pt(size); st.font.bold = True
        st.paragraph_format.line_spacing = Pt(spacing)
        st.paragraph_format.space_before = Pt(9 if name == "Heading 2" else 0)
        st.paragraph_format.space_after = Pt(5 if name == "Heading 2" else 9)
        st.paragraph_format.keep_with_next = True
    st = doc.styles["Caption"]; st.font.size = Pt(8.5); st.font.bold = False; st.font.italic = False
    st.paragraph_format.line_spacing = Pt(12); st.paragraph_format.space_after = Pt(7)
    # No double imposition: the file is already A4 landscape, one side per page.
    for tag in ("bookFoldPrinting", "bookFoldRevPrinting", "mirrorMargins", "printTwoOnOne"):
        for node in doc.settings.element.xpath("./w:"+tag): node.getparent().remove(node)
    doc.core_properties.title = f"Banana Thermal {MODELS[model]['size']}寸单光热成像随附说明书"
    doc.core_properties.author = "Banana Thermal"
    doc.core_properties.subject = f"A4 双面骑马订拼版 {LOGICAL} 面成册 离线使用说明"
    doc.core_properties.comments = ""
    doc.core_properties.keywords = f"standard; {model}; A4; saddle stitch; v1.8"


def compose_body(doc, item, number):
    before = set(doc._element.body)
    pp = para(doc, item["chapter"], size=9, bold=True, color=ORANGE, after=6)
    pp.paragraph_format.keep_with_next = True
    para(doc, item["title"], size=18, bold=True, style="Heading 1")
    pp = para(doc, item["lead"], after=10)
    pp.paragraph_format.keep_with_next = True
    count = 0
    for kind, *values in item["blocks"]:
        if kind == "p": para(doc, values[0])
        elif kind == "h": para(doc, values[0], size=12, bold=True, style="Heading 2")
        elif kind == "n":
            pp = doc.add_paragraph()
            pp.paragraph_format.space_before = Pt(6); pp.paragraph_format.space_after = Pt(5)
            font(pp.add_run(values[0]+"  "), 10.5, True, ORANGE)
            inline(pp, values[1], 10.5)
        elif kind == "s":
            count += 1; pp = doc.add_paragraph()
            pp.paragraph_format.space_before = Pt(5); pp.paragraph_format.space_after = Pt(3)
            pp.paragraph_format.keep_with_next = True
            font(pp.add_run(f"{count:02d}  "), 13, True, ORANGE)
            font(pp.add_run(values[0]), 11.5, True)
            pp = para(doc, values[1], after=6)
            pp.paragraph_format.left_indent = Mm(7)
        elif kind == "i":
            name, caption, width, crop = values
            # Keep large illustrations readable without starving their explanation.
            if "rear-manual" in name: width = 3.15
            if "thermal-view" in name: width = 3.85
            picture(doc, name, caption, min(width, TEXT_W/25.4), crop)
        elif kind == "t": reference_table(doc, *values)
        elif kind == "toc":
            for title, reference in values[0]:
                pp = para(doc, title+"\t"+reference, after=5)
                pp.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W-1), WD_TAB_ALIGNMENT.RIGHT)
        else: raise ValueError(kind)
    return [node for node in doc._element.body if node not in before]


def compose_cover(doc, model):
    before = set(doc._element.body)
    para(doc, "BANANA THERMAL", size=12, bold=True, color=ORANGE, before=7, after=21)
    para(doc, "单光热成像\n使用说明书", size=25, bold=True, style="Title", after=9)
    para(doc, MODELS[model]["size"]+" 寸", size=19, bold=True, after=12)
    picture(doc, "manual-design/cover-standard.png", "外观示意  实际部件见第 2–4 页", 2.85, MODELS[model]["cover"])
    para(doc, "设备操作  测温设置  日常维护", size=11, bold=True, before=10, after=13)
    para(doc, "请在使用前阅读，并随设备妥善保留。", size=10.5, after=4)
    para(doc, "随附说明书 v1.8  ·  2026 年 10 月", size=9, color=GRAY, after=0)
    return [node for node in doc._element.body if node not in before]


def compose_back(doc, model):
    before = set(doc._element.body)
    para(doc, "BANANA THERMAL", size=14, bold=True, color=ORANGE, before=8, after=11)
    para(doc, MODELS[model]["size"]+" 寸单光热成像", size=13, bold=True, after=0)
    para(doc, "操作速记", size=18, bold=True, before=93, after=13)
    para(doc, "{A}  长按换伪彩，短按整理标记。", after=10)
    para(doc, "{B}  短按拍照，再按返回；长按相册。", after=10)
    para(doc, "{A+B}  确认后完全松开两键。", after=10)
    para(doc, "菜单操作与保存见第 7 页。", size=10, color=GRAY, after=0)
    para(doc, "适用统一固件 v1.8.0 及兼容版本。\n旧版迁移前请先备份，见第 25–27 页。", size=9, before=70, after=13)
    para(doc, "产品资料与固件", size=9, bold=True, after=4)
    para(doc, "github.com/applenana/\nBananaThermalFirmware", size=9, color=GRAY, after=12)
    para(doc, "随附说明书 v1.8  ·  2026 年 10 月", size=8.5, color=GRAY, after=0)
    return [node for node in doc._element.body if node not in before]


def furniture(doc, model, item, number, footer=False):
    before = set(doc._element.body)
    if footer:
        pp = para(doc, item["chapter"]+f"\t{number:02d}", size=8.5, color=GRAY, after=0)
    else:
        pp = para(doc, "BANANA THERMAL\t"+MODELS[model]["size"]+" 寸单光", size=8, after=0)
    pp.paragraph_format.line_spacing = Pt(12)
    pp.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W-1), WD_TAB_ALIGNMENT.RIGHT)
    return [node for node in doc._element.body if node not in before]


def box(anchor, nodes, x, y, width, height, shape_id):
    """Native Word VML page frame, not a rasterized page or a narrative callout."""
    pict = OxmlElement("w:pict")
    shape = etree.SubElement(pict, "{"+NS_V+"}shape", nsmap={"v":NS_V,"o":NS_O,"w10":NS_W10})
    shape.set("id", f"_x0000_s{shape_id}"); shape.set("type", "#_x0000_t202")
    shape.set("filled", "f"); shape.set("stroked", "f")
    shape.set("{"+NS_O+"}allowincell", "f")
    pt = lambda mm: mm * 72/25.4
    shape.set("style", f"position:absolute;margin-left:{pt(x):.4f}pt;margin-top:{pt(y):.4f}pt;"
              f"width:{pt(width):.4f}pt;height:{pt(height):.4f}pt;z-index:{shape_id};"
              "mso-position-horizontal-relative:page;mso-position-vertical-relative:page;"
              "mso-wrap-style:square;v-text-anchor:top")
    textbox = etree.SubElement(shape, "{"+NS_V+"}textbox")
    textbox.set("inset", "0,0,0,0"); textbox.set("style", "mso-fit-shape-to-text:f")
    content = etree.SubElement(textbox, qn("w:txbxContent"))
    for node in nodes: content.append(node)
    if not nodes or nodes[-1].tag != qn("w:p"): content.append(OxmlElement("w:p"))
    etree.SubElement(shape, "{"+NS_W10+"}wrap").set("type", "none")
    anchor.add_run()._r.append(pict)


def build(model):
    doc = Document(); setup(doc, model); plan = book(model)
    mapping = pairs(); shape_id = 1000
    for side_no, pair in enumerate(mapping):
        anchor = doc.add_paragraph()
        pf = anchor.paragraph_format
        pf.page_break_before = side_no > 0
        pf.space_before = pf.space_after = Pt(0); pf.line_spacing = Pt(1)
        for half, logical in enumerate(pair):
            x = half * HALF + (14 if half == 0 else 15)
            if logical in (2, LOGICAL-1): continue
            if logical == 1:
                nodes = compose_cover(doc, model); y, height = 16, 180
            elif logical == LOGICAL:
                nodes = compose_back(doc, model); y, height = 16, 180
            else:
                number = logical-2; item = plan[number-1]
                shape_id += 1; box(anchor, furniture(doc, model, item, number), x, 9, TEXT_W, 7, shape_id)
                shape_id += 1; box(anchor, furniture(doc, model, item, number, True), x, 196, TEXT_W, 6, shape_id)
                nodes = compose_body(doc, item, number); y, height = 22, 169
            shape_id += 1; box(anchor, nodes, x, y, TEXT_W, height, shape_id)
    OUTPUT.mkdir(parents=True, exist_ok=True)
    output = OUTPUT / (stem(model)+".docx")
    if output.exists():
        raise FileExistsError(f"Preserving existing DOCX; choose a new filename: {output}")
    doc.save(output)
    work = ROOT / ".work/print-booklets" / model; work.mkdir(parents=True, exist_ok=True)
    (work / "imposition.json").write_text(json.dumps(dict(model=model, logical_pages=LOGICAL,
        body_pages=BODY_PAGES, sheets=LOGICAL//4, sides=mapping), ensure_ascii=False, indent=2), encoding="utf-8")
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--model", choices=["all", *MODELS], default="all")
    args = parser.parse_args()
    for model in MODELS if args.model == "all" else [args.model]: print(build(model))
