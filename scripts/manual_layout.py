"""Native Word page system for the two product manuals."""
from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from manual_content import pages

ROOT = Path(__file__).resolve().parents[1]
MANUALS = {
    "standard": ("single-light.md", "BananaThermal_单光热成像使用说明书_v1.8"),
    "dual": ("dual-light.md", "BananaThermal_3.2寸双光热成像使用说明书_v1.8"),
}
REPO = "https://github.com/applenana/BananaThermalFirmware/blob/main/"
ORANGE, NAVY, GRAY = "D94B17", "183D50", "5B6570"
WIDTH = 6.98


def el(name, **attrs):
    node = OxmlElement(name)
    for key, value in attrs.items(): node.set(qn(key) if ":" in key else key, str(value))
    return node


def font(run, size=None, bold=None, color="000000"):
    run.font.name = "Calibri"
    rp = run._element.get_or_add_rPr()
    rp.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    return run


def inline(para, text, size=None):
    for part in re.split(r"(\*\*.*?\*\*|\{[^}]+\})", text):
        if not part: continue
        if part.startswith("{"):
            r = font(para.add_run(" " + part[1:-1] + " "), size or 10.5, True, "FFFFFF")
            r._element.get_or_add_rPr().append(el("w:shd", **{"w:fill": NAVY}))
        elif part.startswith("**"):
            font(para.add_run(part[2:-2]), size, True)
        else:
            font(para.add_run(part.replace("_", "_\u200b")), size)
    return para


def body(doc, text, *, before=0, after=8, size=None):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    inline(para, text, size)
    return para


def heading(doc, text):
    para = doc.add_paragraph(style="Heading 2")
    font(para.add_run(text), 13, True)
    return para


def bookmark(para, name, number):
    para._p.insert(0, el("w:bookmarkStart", **{"w:id": number, "w:name": name}))
    para._p.append(el("w:bookmarkEnd", **{"w:id": number}))


def link(para, text, target, internal=False, size=11, color=NAVY):
    node = el("w:hyperlink")
    if internal: node.set(qn("w:anchor"), target)
    else:
        node.set(qn("r:id"), para.part.relate_to(target,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True))
    node.set(qn("w:history"), "1")
    r = el("w:r"); prop = el("w:rPr")
    prop.append(el("w:rFonts", **{"w:ascii": "Calibri", "w:hAnsi": "Calibri", "w:eastAsia": "Microsoft YaHei"}))
    prop.append(el("w:color", **{"w:val": color}))
    prop.append(el("w:sz", **{"w:val": int(size*2)}))
    prop.append(el("w:u", **{"w:val": "single"}))
    r.append(prop)
    txt = el("w:t"); txt.text = text; r.append(txt); node.append(r); para._p.append(node)


def picture(doc, name, caption, width=6.7, crop=None):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(4)
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1
    para.paragraph_format.keep_with_next = True
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = para.add_run().add_picture(str(ROOT / "docs/assets" / name), width=Inches(width))
    if crop:
        # Native Word crop: preserve the source image and avoid distorting it.
        left, top, right, bottom = crop
        from PIL import Image
        with Image.open(ROOT / "docs/assets" / name) as image:
            ratio = image.height * (1-top-bottom) / (image.width * (1-left-right))
        shape.height = Inches(width * ratio)
        fill = shape._inline.xpath(".//pic:blipFill")[0]
        fill.insert(1, el("a:srcRect", l=int(left*100000), t=int(top*100000),
                          r=int(right*100000), b=int(bottom*100000)))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption.split("  ")[0])
    para = doc.add_paragraph(style="Caption")
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(para.add_run(caption), 9, False, GRAY)


def table(doc, headers, rows, widths):
    tb = doc.add_table(rows=0, cols=len(headers)); tb.autofit = False
    widths = [w * WIDTH / sum(widths) for w in widths]
    for col, width in zip(tb.columns, widths): col.width = Inches(width)
    pr = tb._tbl.tblPr
    borders = el("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(el("w:"+side, **{"w:val": "single", "w:sz": 4, "w:color": "D9D9D9"}))
    pr.append(borders)
    margins = el("w:tblCellMar")
    for side, value in (("top", 100), ("bottom", 100), ("left", 125), ("right", 125)):
        margins.append(el("w:"+side, **{"w:w": value, "w:type": "dxa"}))
    pr.append(margins)
    for index, values in enumerate([headers, *rows]):
        row = tb.add_row(); trpr = row._tr.get_or_add_trPr()
        trpr.append(el("w:cantSplit"))
        if index == 0: trpr.append(el("w:tblHeader"))
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell._tc.get_or_add_tcPr().append(el("w:shd", **{"w:fill": NAVY if index == 0 else ("F2F5F7" if index % 2 == 0 else "FFFFFF")}))
            para = cell.paragraphs[0]; pf = para.paragraph_format
            pf.line_spacing = Pt(15); pf.space_after = Pt(0); pf.space_before = Pt(0)
            pf.keep_with_next = index == 0
            if headers[j] in ("Flash", "LittleFS", "CPU", "原生阵列", "原生", "全屏热流", "方屏热流", "初始值", "进入的菜单"):
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            font(para.add_run(value.replace("_", "_\u200b")), 10, index == 0, "FFFFFF" if index == 0 else "000000")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0); spacer.paragraph_format.line_spacing = Pt(6)


def field(para, command, cached="1"):
    run = para.add_run()
    run._r.append(el("w:fldChar", **{"w:fldCharType": "begin"}))
    node = el("w:instrText"); node.text = command; run._r.append(node)
    run._r.append(el("w:fldChar", **{"w:fldCharType": "separate"}))
    txt = el("w:t"); txt.text = cached; run._r.append(txt)
    run._r.append(el("w:fldChar", **{"w:fldCharType": "end"}))
    font(run, 9.5, color=GRAY)


def setup(doc, profile):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(.76)
    sec.top_margin = sec.bottom_margin = Inches(.68)
    sec.header_distance = Inches(.28); sec.footer_distance = Inches(.3)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption", "Header", "Footer"):
        style = doc.styles[name]; style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        for node in style._element.xpath(".//w:color"):
            for key in ("w:themeColor", "w:themeTint", "w:themeShade"): node.attrib.pop(qn(key), None)
        for node in style._element.xpath(".//w:pBdr"): node.getparent().remove(node)
    normal = doc.styles["Normal"]; normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = Pt(17); normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.widow_control = True
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        pf = doc.styles[name].paragraph_format; pf.keep_with_next = True; pf.keep_together = True
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(12)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(6)
    doc.styles["Heading 2"].paragraph_format.line_spacing = Pt(20)
    doc.styles["Caption"].paragraph_format.space_after = Pt(10)
    doc.styles["Caption"].paragraph_format.line_spacing = Pt(13)
    doc.styles["Caption"].font.bold = False
    doc.styles["Caption"].font.italic = False
    for name in ("Header", "Footer"):
        # Override the built-in centered tab at 3.25 inches as well as the right tab.
        style = doc.styles[name]
        style.paragraph_format.tab_stops.clear_all()
        for pos in (4680, 9360):
            style.element.get_or_add_pPr().get_or_add_tabs().append(el("w:tab", **{"w:val": "clear", "w:pos": pos}))
        style.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    doc.settings.element.append(el("w:updateFields", **{"w:val": "true"}))
    doc.core_properties.title = "Banana Thermal " + ("单光" if profile == "standard" else "双光") + "热成像使用说明书"
    doc.core_properties.author = "Banana Thermal"
    doc.core_properties.subject = "设备操作、测温、数据保存与固件升级"
    doc.core_properties.keywords = "v1.8; " + profile
    doc.core_properties.comments = ""


def start_page(doc, profile, page_no, title, chapter):
    sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    last = doc.paragraphs[-1]
    last.paragraph_format.line_spacing = Pt(1); last.paragraph_format.space_after = Pt(0)
    last.paragraph_format.space_before = Pt(0)
    for story in (sec.header, sec.footer): story.is_linked_to_previous = False
    hp = sec.header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    hp.paragraph_format.space_after = Pt(0)
    font(hp.add_run("BANANA THERMAL"), 8.5, True)
    font(hp.add_run("\t" + ("单光" if profile == "standard" else "双光") + "热成像使用说明书"), 8.5)
    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    fp.paragraph_format.space_after = Pt(0)
    font(fp.add_run("v1.8  ·  " + chapter), 9, color=GRAY); font(fp.add_run("\t"), 9)
    field(fp, "PAGE", str(page_no)); font(fp.add_run(" / "), 9.5, color=GRAY)
    field(fp, "NUMPAGES", "20")
    ey = doc.add_paragraph()
    ey.paragraph_format.space_before = Pt(8); ey.paragraph_format.space_after = Pt(8)
    font(ey.add_run(f"{page_no:02d}  /  {chapter}"), 9.5, True, ORANGE)
    titlep = doc.add_paragraph(style="Heading 1")
    titlep.paragraph_format.line_spacing = Pt(32)
    font(titlep.add_run(title), 24, True); bookmark(titlep, f"page_{page_no}", page_no)


def cover(doc, profile):
    standard = profile == "standard"
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22); p.paragraph_format.space_after = Pt(26)
    font(p.add_run("BANANA THERMAL"), 15, True, ORANGE)
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(14); p.paragraph_format.line_spacing = Pt(48)
    font(p.add_run(("单光" if standard else "双光") + "热成像\n使用说明书"), 34, True)
    body(doc, "2.4 寸  /  2.8 寸  /  3.2 寸" if standard else "3.2 寸双光融合设备", after=6, size=14)
    body(doc, "STANDARD" if standard else "DUAL", after=22, size=10)
    picture(doc, "manual-design/cover-" + profile + ".png", "设备外形与热像示意  具体外壳、传感器与接口以实物为准", WIDTH)
    p = body(doc, "从开机到测温 从照片到数据", before=18, after=10, size=15)
    for r in p.runs: r.bold = True
    body(doc, "本手册帮助你完成设备操作、参数设置、温度校准与报警、照片管理、Studio 连接和固件升级。" + ("三种单光机型共用 standard 固件，按键和外壳差异在各自页面说明。" if standard else "两路画面对齐与可选测距单独说明，便于按实物配置操作。"), after=20)
    body(doc, "文档 v1.8   ·   2026 年 10 月", after=4, size=10)
    body(doc, "固件行为基线 v1.8.0   /   Studio 0.5.5", after=0, size=9.5)


def contents(doc, profile, plan):
    start_page(doc, profile, 2, "按任务阅读", "阅读导航")
    body(doc, "第一次使用从第 3 页开始；需要解决具体问题时，点击下面的页名直接跳转。", after=14)
    groups = [("开始使用", 3, 3), ("设备操作", 4, 10), ("测温与报警", 11, 13),
              ("照片与 Studio", 14, 16), ("升级与维护", 17, 20)]
    for label, first, last in groups:
        pp = doc.add_paragraph()
        pp.paragraph_format.space_before = Pt(9); pp.paragraph_format.space_after = Pt(3)
        pp.paragraph_format.line_spacing = Pt(17); pp.paragraph_format.keep_with_next = True
        font(pp.add_run(label), 11, True)
        for num in range(first, last + 1):
            pp = doc.add_paragraph()
            pp.paragraph_format.space_after = Pt(1); pp.paragraph_format.line_spacing = Pt(20)
            pp.paragraph_format.left_indent = Inches(.17)
            pp.paragraph_format.tab_stops.add_tab_stop(Inches(WIDTH-.17), WD_TAB_ALIGNMENT.RIGHT)
            link(pp, plan[num-1]["title"], f"page_{num}", internal=True)
            font(pp.add_run(f"\t{num:02d}"), 11, color=GRAY)
    body(doc, "使用前确认板型与 Flash；刷机前先备份照片、配置和校准。", before=13, size=10)


def build(profile):
    doc = Document(); setup(doc, profile); plan = pages(profile)
    cover(doc, profile); contents(doc, profile, plan)
    for num, item in enumerate(plan[2:], 3):
        start_page(doc, profile, num, item["title"], item["chapter"])
        lead = body(doc, item["lead"], after=14, size=12)
        lead.paragraph_format.line_spacing = Pt(19); lead.paragraph_format.keep_with_next = True
        count = 0
        for kind, *values in item["blocks"]:
            if kind == "p": body(doc, values[0])
            elif kind == "h": heading(doc, values[0])
            elif kind == "n":
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(9); p.paragraph_format.space_after = Pt(6)
                font(p.add_run(values[0] + "  "), 10.5, True, ORANGE); inline(p, values[1], 10.5)
            elif kind == "s":
                count += 1; pp = doc.add_paragraph()
                pp.paragraph_format.space_before = Pt(9); pp.paragraph_format.space_after = Pt(4)
                pp.paragraph_format.keep_with_next = True
                font(pp.add_run(f"{count:02d}  "), 16, True, ORANGE)
                font(pp.add_run(values[0]), 12, True)
                pp = body(doc, values[1], after=8); pp.paragraph_format.left_indent = Inches(.38)
            elif kind == "i": picture(doc, *values)
            elif kind == "t": table(doc, *values)
            elif kind == "l":
                label, target = values; pp = doc.add_paragraph()
                pp.paragraph_format.space_after = Pt(3); pp.paragraph_format.line_spacing = Pt(19)
                link(pp, label, target if target.startswith("http") else REPO+target)
    output = ROOT / "artifacts" / (MANUALS[profile][1] + ".docx")
    doc.save(output); return output
