#!/usr/bin/env python3
"""Historical v1.7 dual-light manual builder.

Current single-light and dual-light manuals use scripts/build_manuals.py.
Keep this builder only to reproduce the archived v1.7 documentation.

The document follows the compact_reference_guide preset from the bundled
documents skill: US Letter, 1 inch margins, 11 pt body, 1.25 line spacing,
real Word headings/numbering, fixed-DXA table geometry, and restrained blue
accents. Chinese text uses Microsoft YaHei as an explicit CJK fallback.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts" / "BananaThermal_3.2寸双光热成像使用说明书_v1.7.docx"

ASSETS = ROOT / "docs" / "assets" / "user-guide"
HARDWARE_IMAGE = ASSETS / "device-controls.png"
DUAL_IMAGE = ASSETS / "dual-light-principles.png"
STUDIO_REALTIME = ASSETS / "studio-realtime.png"
STUDIO_DEVICE_GALLERY = ASSETS / "studio-device-gallery.png"
STUDIO_SOFTWARE_GALLERY = ASSETS / "studio-software-gallery.png"

# compact_reference_guide exact tokens
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "203748"
INK = "243241"
MUTED = "667788"
PALE_BLUE = "E8EEF5"
PALE_GRAY = "F2F4F7"
PALE_TEAL = "E7F6F4"
PALE_GOLD = "FFF5D9"
PALE_RED = "FCEBEA"
BORDER = "CBD6E2"
WHITE = "FFFFFF"
ACCENT_ORANGE = "F0642A"
CAUTION = "7A5A00"
RISK = "9B1C1C"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(
    run,
    *,
    latin: str = "Calibri",
    east_asia: str = "Microsoft YaHei",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), latin)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, margins: dict[str, int] = CELL_MARGINS) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "bottom", "start", "end"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margins[side]))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = BORDER, size: int = 6) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def apply_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {CONTENT_WIDTH_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for col_idx, width in enumerate(widths):
        table.columns[col_idx].width = Twips(width)
    for row in table.rows:
        row.height = None
        for col_idx, cell in enumerate(row.cells):
            width = widths[col_idx]
            cell.width = Twips(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            set_cell_margins(cell)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_table_row_together(row) -> None:
    """Prevent Word from splitting a logical table row across two pages."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_paragraph_left_border(paragraph, color: str, size: int = 18, space: int = 12) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = p_bdr.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        p_bdr.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)


def set_alt_text(inline_shape, title: str, description: str) -> None:
    inline_shape._inline.docPr.set("title", title)
    inline_shape._inline.docPr.set("descr", description)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


def add_numbering_definitions(doc: Document) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element
    nsid_seed = 900
    next_abstract = max(
        [int(n.get(qn("w:abstractNumId"))) for n in numbering.findall(qn("w:abstractNum"))]
        + [0]
    ) + 1
    next_num = max(
        [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))] + [0]
    ) + 1

    def build(abstract_id: int, num_id: int, fmt: str, text: str, font: str) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        nsid = OxmlElement("w:nsid")
        nsid.set(qn("w:val"), f"{nsid_seed + abstract_id:08X}"[-8:])
        abstract.append(nsid)
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        p_pr.extend([tabs, ind])
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), font)
        r_fonts.set(qn("w:hAnsi"), font)
        r_fonts.set(qn("w:eastAsia"), font)
        r_pr.append(r_fonts)
        lvl.extend([start, num_fmt, lvl_text, suff, lvl_jc, p_pr, r_pr])
        abstract.append(lvl)
        numbering.append(abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_id_el = OxmlElement("w:abstractNumId")
        abstract_id_el.set(qn("w:val"), str(abstract_id))
        num.append(abstract_id_el)
        numbering.append(num)

    build(next_abstract, next_num, "bullet", "•", "Arial")
    build(next_abstract + 1, next_num + 1, "decimal", "%1.", "Arial")
    return next_num, next_num + 1


def assign_num(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.insert(0, num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])


def restart_numbering(doc: Document, base_num_id: int) -> int:
    """Create a new concrete numbering instance so a procedure restarts at 1."""
    numbering = doc.part.numbering_part.element
    base = next(
        node
        for node in numbering.findall(qn("w:num"))
        if int(node.get(qn("w:numId"))) == base_num_id
    )
    abstract_id = base.find(qn("w:abstractNumId")).get(qn("w:val"))
    next_num_id = max(
        [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))] + [0]
    ) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(next_num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.extend([abstract, override])
    numbering.append(num)
    return next_num_id


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.25
    pf.widow_control = True

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    for name in ("Manual Bullet", "Manual Number"):
        if name not in styles:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[name]
        style.base_style = normal
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    if "Table Body" not in styles:
        table_body = styles.add_style("Table Body", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_body = styles["Table Body"]
    table_body.base_style = normal
    table_body.font.name = "Calibri"
    table_body.font.size = Pt(9.5)
    table_body.font.color.rgb = rgb(INK)
    table_body._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    table_body.paragraph_format.space_before = Pt(0)
    table_body.paragraph_format.space_after = Pt(0)
    table_body.paragraph_format.line_spacing = 1.15

    if "Caption Manual" not in styles:
        caption = styles.add_style("Caption Manual", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Caption Manual"]
    caption.base_style = normal
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = rgb(MUTED)
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = False


def configure_document(doc: Document) -> tuple[int, int]:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_ALIGN_PARAGRAPH.RIGHT)
    r = p.add_run("BANANA THERMAL · 使用说明书")
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    r = p.add_run("\t固件 v1.7.0 / Studio v0.4.2")
    set_run_font(r, size=8.5, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    r = fp.add_run("Banana Thermal  ·  ")
    set_run_font(r, size=8.5, color=MUTED)
    add_field(fp, "PAGE", "1")

    first_footer = section.first_page_footer
    ffp = first_footer.paragraphs[0]
    ffp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = ffp.add_run("用户手册 · 2026-07-31 · 请与对应版本发布说明配合使用")
    set_run_font(r, size=8.5, color=MUTED)

    configure_styles(doc)

    # Reuse Word's built-in list definitions.  Custom numbering definitions can
    # be normalized incorrectly when desktop Word updates the TOC and saves the
    # package, which may turn bullet paragraphs into a continued decimal list.
    # The built-in definitions survive that round-trip reliably; numbered
    # procedures still receive fresh concrete instances in ``number()``.
    bullet_num_id = int(
        doc.styles["List Bullet"]._element.pPr.numPr.numId.val
    )
    decimal_num_id = int(
        doc.styles["List Number"]._element.pPr.numPr.numId.val
    )
    return bullet_num_id, decimal_num_id


class ManualBuilder:
    def __init__(self, doc: Document, bullet_num_id: int, decimal_num_id: int):
        self.doc = doc
        self.bullet_num_id = bullet_num_id
        self.decimal_num_id = decimal_num_id
        self.figure_number = 0
        self._number_active = False
        self._active_decimal_num_id = decimal_num_id

    def body(self, text: str, *, bold_lead: str | None = None) -> None:
        self._number_active = False
        p = self.doc.add_paragraph(style="Normal")
        if bold_lead and text.startswith(bold_lead):
            r = p.add_run(bold_lead)
            set_run_font(r, bold=True)
            r = p.add_run(text[len(bold_lead):])
            set_run_font(r)
        else:
            r = p.add_run(text)
            set_run_font(r)

    def bullet(self, text: str) -> None:
        self._number_active = False
        p = self.doc.add_paragraph(style="Manual Bullet")
        assign_num(p, self.bullet_num_id)
        r = p.add_run(text)
        set_run_font(r)

    def number(self, text: str) -> None:
        if not self._number_active:
            self._active_decimal_num_id = restart_numbering(self.doc, self.decimal_num_id)
            self._number_active = True
        p = self.doc.add_paragraph(style="Manual Number")
        assign_num(p, self._active_decimal_num_id)
        r = p.add_run(text)
        set_run_font(r)

    def h1(self, text: str, *, page_break: bool = False) -> None:
        self._number_active = False
        p = self.doc.add_paragraph(text, style="Heading 1")
        p.paragraph_format.page_break_before = page_break
        p.paragraph_format.left_indent = Inches(0)
        p.paragraph_format.right_indent = Inches(0)
        p.paragraph_format.first_line_indent = Inches(0)
        for run in p.runs:
            set_run_font(run, size=16, color=BLUE, bold=True)

    def h2(self, text: str, *, page_break: bool = False) -> None:
        self._number_active = False
        p = self.doc.add_paragraph(text, style="Heading 2")
        p.paragraph_format.page_break_before = page_break
        p.paragraph_format.left_indent = Inches(0)
        p.paragraph_format.right_indent = Inches(0)
        p.paragraph_format.first_line_indent = Inches(0)
        for run in p.runs:
            set_run_font(run, size=13, color=BLUE, bold=True)

    def h3(self, text: str) -> None:
        self._number_active = False
        p = self.doc.add_paragraph(text, style="Heading 3")
        p.paragraph_format.left_indent = Inches(0)
        p.paragraph_format.right_indent = Inches(0)
        p.paragraph_format.first_line_indent = Inches(0)
        for run in p.runs:
            set_run_font(run, size=12, color=DARK_BLUE, bold=True)

    def callout(self, label: str, text: str, kind: str = "info") -> None:
        self._number_active = False
        colors = {
            "info": (PALE_BLUE, BLUE),
            "tip": (PALE_TEAL, "178C7A"),
            "warn": (PALE_GOLD, CAUTION),
            "risk": (PALE_RED, RISK),
        }
        fill, accent = colors[kind]
        p = self.doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.right_indent = Inches(0.08)
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.2
        set_paragraph_shading(p, fill)
        set_paragraph_left_border(p, accent)
        r = p.add_run(f"{label}  ")
        set_run_font(r, bold=True, color=accent)
        r = p.add_run(text)
        set_run_font(r, color=INK)

    def table(
        self,
        headers: list[str],
        rows: list[list[str]],
        widths: list[int],
        *,
        compact: bool = False,
    ) -> None:
        self._number_active = False
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        set_table_borders(table)
        hdr = table.rows[0]
        repeat_table_header(hdr)
        keep_table_row_together(hdr)
        for idx, text in enumerate(headers):
            cell = hdr.cells[idx]
            set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.style = self.doc.styles["Table Body"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            r = p.add_run(text)
            set_run_font(r, size=9.2 if compact else 9.5, color=NAVY, bold=True)
        for row_index, values in enumerate(rows):
            row = table.add_row()
            keep_table_row_together(row)
            cells = row.cells
            if row_index % 2 == 1:
                for cell in cells:
                    set_cell_shading(cell, "FAFBFC")
            for idx, text in enumerate(values):
                p = cells[idx].paragraphs[0]
                p.style = self.doc.styles["Table Body"]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.keep_together = True
                if row_index == 0 and idx < len(values) - 1:
                    p.paragraph_format.keep_with_next = True
                r = p.add_run(str(text))
                set_run_font(r, size=9 if compact else 9.5, color=INK)
        apply_table_geometry(table, widths)
        after = self.doc.add_paragraph()
        after.paragraph_format.space_before = Pt(0)
        after.paragraph_format.space_after = Pt(2)

    def figure(self, path: Path, caption: str, width: float, alt: str) -> None:
        self._number_active = False
        if not path.exists():
            self.callout("图片未找到", f"生成时未找到 {path.name}，请检查资料路径。", "warn")
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        shape = p.add_run().add_picture(str(path), width=Inches(width))
        set_alt_text(shape, caption, alt)
        self.figure_number += 1
        cp = self.doc.add_paragraph(style="Caption Manual")
        r = cp.add_run(f"图 {self.figure_number}  {caption}")
        set_run_font(r, size=9, color=MUTED, italic=True)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BANANA THERMAL")
    set_run_font(r, size=11, color=ACCENT_ORANGE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("3.2 寸双光融合热成像")
    set_run_font(r, size=28, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("设备操作 · 菜单配置 · 上位机使用 · 固件升级")
    set_run_font(r, size=13.5, color=MUTED)

    if HARDWARE_IMAGE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(16)
        shape = p.add_run().add_picture(str(HARDWARE_IMAGE), width=Inches(5.7))
        set_alt_text(shape, "设备结构总览", "3.2 寸双光热成像设备的外观、按键、接口和传感器位置示意")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("适用版本")
    set_run_font(r, size=9.5, color=MUTED, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("固件 v1.7.0  ·  Next BananaThermal Studio v0.4.2")
    set_run_font(r, size=10.5, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("依据当前发布文档、设备源码、上位机源码与硬件线稿整理")
    set_run_font(r, size=9, color=MUTED, italic=True)


def add_front_matter(b: ManualBuilder) -> None:
    b.h1("使用前必读", page_break=True)
    b.callout(
        "用途边界",
        "本设备适合观察温度分布、发现异常发热点和进行趋势对比；不属于医疗体温计、消防认证仪器或计量级温度标准器。涉及人身、消防、电气安全或合规判定时，请使用相应认证设备复核。",
        "risk",
    )
    b.body("本说明书优先回答三个问题：设备上的按钮怎么按、菜单参数怎么调、上位机里的每个常用操作怎么完成。功能会随传感器、摄像头、距离模块和屏幕批次自动增减，因此您看到的菜单可能与插图略有差异。")
    b.table(
        ["标识", "含义", "建议动作"],
        [
            ["提示", "能提高效率或画面质量的经验", "可按场景选择"],
            ["注意", "错误设置可能造成读数或显示异常", "先记录原值，再逐项修改"],
            ["危险", "可能造成数据丢失、刷机失败或设备损坏", "确认条件满足后再执行"],
        ],
        [1200, 3600, 4560],
    )
    b.h2("版本与硬件差异")
    b.bullet("热阵列可能是 MLX90640 / MLX90641（物理 32×24）或 Heimann HTPA32x32（物理 32×32）。")
    b.bullet("摄像头可能是 OV2640 或 OV3660；没有摄像头或初始化失败时，摄像头与融合相关项目不可用。")
    b.bullet("VL53L0X 距离传感器为可选件；只有初始化成功时，菜单才显示“视差调整”和“距离传感器”。")
    b.bullet("部分旧款机型没有独立左/右/确认拨片，可用按钮 A、按钮 B 及 A+B 组合完成菜单操作，详见“按键总览”。")
    b.callout("版本说明", "本文以固件 v1.7.0 与 Studio v0.4.2 的当前源码为准。升级后若菜单名称、默认值或范围改变，请同时查看对应 Release Notes。", "info")

    b.h1("目录", page_break=True)
    p = b.doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "在 Word 中右键此处并选择“更新域”即可刷新目录。")
    b.callout("阅读建议", "初次使用按“设备与按键 → 5 分钟快速上手 → 上位机连接”阅读；需要调画质或对齐时再查 16 项菜单；遇到异常直接跳到“故障排查”。", "tip")


def add_device_overview(b: ManualBuilder) -> None:
    b.h1("1 认识设备", page_break=True)
    b.body("设备把热阵列测得的温度分布与可见光摄像头的轮廓信息组合到同一画面。热像负责“哪里热、热多少”，可见光负责“这是什么、位置在哪里”；两路传感器物理位置不同，所以近距离使用时需要做视差补偿。")
    b.figure(
        HARDWARE_IMAGE,
        "设备结构、按键、接口与传感器位置示意",
        6.2,
        "设备正面、顶部、底部和背面示意；标出开关、菜单拨片、按钮 A、按钮 B、复位键、快捷菜单键、Type-C 口、充电灯及三类传感器。",
    )
    b.callout("以实物为准", "传感器位置、外壳细节和是否配有测距模块会因批次不同而变化；不要仅凭外观判断热阵列型号。", "warn")
    b.h2("1.1 主要部件")
    b.table(
        ["部件", "作用", "使用要点"],
        [
            ["电源开关", "控制整机电源", "拨到开机位置后等待传感器初始化；关机前先结束写入或下载。"],
            ["菜单拨片 / 左右确认键", "进入菜单、移动选项、编辑数值", "按下用于进入或确认；左右拨动用于选择或增减。"],
            ["按钮 A", "清理测温标记、切换色板、相册删除", "菜单中可充当左键；长按约 1 秒切换伪彩。"],
            ["按钮 B", "拍照、退出拍照状态、进入相册", "菜单中可充当右键；长按约 1 秒进入设备相册。"],
            ["快捷菜单键", "开启或关闭悬浮快捷菜单", "开启后用左右键浏览或调值，用确认键切换浏览/编辑。"],
            ["复位键", "立即重启主处理器和传感器", "不会代替正常关机；刷机时与按钮 A 配合进入 UF2 模式。"],
            ["Type-C 接口", "供电、充电与 USB 串口通讯", "使用质量可靠的数据线；仅充电线无法连接上位机。"],
            ["充电指示灯", "显示充电状态", "常见状态为充电亮、充满灭；不同充电板批次可能略有差异。"],
        ],
        [1500, 3000, 4860],
        compact=True,
    )
    b.h2("1.2 热成像数据不是屏幕像素")
    b.table(
        ["模式 / 传感器", "真实温度阵列", "设备显示 / 上位机说明"],
        [
            ["全屏 · 任意传感器", "32×24", "LCD 为 320×240；上位机按 768 个 float32 温度点重建。"],
            ["方屏 · MLX", "32×24", "LCD 为 240×240；软件按方屏视图参数复原，不伪造额外温度行。"],
            ["方屏 · Heimann", "32×32", "LCD 为 240×240；上位机接收 1024 个 float32 温度点。"],
        ],
        [2200, 2200, 4960],
    )
    b.h2("1.3 主画面状态栏")
    b.body("状态栏把最常用的运行信息集中在画面边缘。字段会随硬件能力和当前模式变化；数值短暂跳动不一定代表故障，应结合实时画面与日志判断。")
    b.table(
        ["字段", "含义", "怎么看"],
        [
            ["最高温 / 最低温", "当前可见热像区域内的极值", "缩放、平移、方屏裁剪或旋转后，统计区域也会随之变化。"],
            ["融合", "当前双光融合模式", "关闭、混合、边缘或半采样；没有摄像头时相关状态不可用。"],
            ["电量", "设备估算的电池百分比", "充电时仍可能缓慢变化；异常跳变先检查供电与电池连接。"],
            ["耗时", "单帧渲染/处理用时，单位 ms", "数值持续升高通常表示滤波、融合或推流负载较重。"],
            ["距离", "VL53L0X 测得或固件判定的目标距离", "N/A 表示模块不可用；>2m 或 >3m 表示超出对应工作区间，不是精确距离。"],
        ],
        [1600, 3400, 4360],
        compact=True,
    )


def add_buttons_and_quick_start(b: ManualBuilder) -> None:
    b.h1("2 按键总览与 5 分钟快速上手", page_break=True)
    b.h2("2.1 不同界面下的按键")
    b.table(
        ["按键", "主画面", "菜单 / 快捷菜单", "设备相册"],
        [
            ["左拨 / 左键", "快捷菜单关闭时，循环降低二次滤波；开启时浏览上一项或减小数值", "上一项；编辑时减小。长按会逐渐加速重复", "上一张"],
            ["按下拨片 / 确认", "打开详细菜单；快捷菜单开启时进入/退出编辑", "进入栏目、切换编辑状态、确认；选择“返回”时保存并返回", "取消删除提示或执行当前确认"],
            ["右拨 / 右键", "快捷菜单关闭时，循环提高二次滤波；开启时浏览下一项或增大数值", "下一项；编辑时增大。长按会逐渐加速重复", "下一张"],
            ["按钮 A 短按", "依次清除多点、隐藏主光标、隐藏冷热追踪；再次按恢复", "旧款无拨片机型中充当左键", "第一次弹出删除确认，第二次确认删除"],
            ["按钮 A 长按", "循环切换 7 套伪彩", "菜单内长按按键用于连续左移", "保持当前相册行为"],
            ["按钮 B 短按", "拍摄并冻结当前画面；再次按退出拍照状态", "旧款无拨片机型中充当右键", "退出相册"],
            ["按钮 B 长按", "进入设备相册并加载照片", "菜单内长按按键用于连续右移", "进入/保持照片预览"],
            ["A + B 同时按", "旧款无独立确认键时等同“确认”", "进入、确认或切换编辑", "取消当前删除流程"],
            ["快捷菜单键", "开启/关闭快捷菜单", "关闭时保存状态；修改后约 3 秒自动保存", "无主要操作"],
        ],
        [1300, 2850, 3060, 2150],
        compact=True,
    )
    b.callout("相册删除不可恢复", "在设备相册中，按钮 A 需要连续确认两次才删除当前图片；任何其他按键会取消。删除后无法从设备恢复，重要照片请先用上位机下载。", "risk")

    b.h2("2.2 快捷菜单怎么用")
    b.number("按“快捷菜单键”，屏幕提示“快捷菜单：开启”。此时默认是浏览状态。")
    b.number("左右拨动选择条目；按下拨片进入编辑，数值两侧出现强调符号。")
    b.number("在编辑状态左右拨动改变数值；再次按下拨片退出编辑。停止操作约 3 秒后自动保存。")
    b.table(
        ["组别", "快捷条目"],
        [
            ["高频显示", "全屏、亮度、色板、可见光缩放"],
            ["距离对齐", "视差模式、自动基准、可见光 X、可见光 Y"],
            ["像素对齐", "热像缩放、热像 X、热像 Y"],
            ["融合效果", "融合模式、混合比例、伽马、边缘颜色"],
        ],
        [1900, 7460],
    )
    b.callout("项目会自动隐藏", "没有摄像头时，对应值显示 N/A；没有 VL53L0X 时，“视差模式”和“自动基准”不会出现在快捷菜单中。", "info")
    b.callout("亮度范围有一处差异", "详细菜单允许 0–256；快捷菜单为避免误触造成近似黑屏，会把亮度限制在 5–256。需要完全熄灭背光时请用详细菜单。", "tip")

    b.h2("2.3 第一次开机")
    for step in [
        "用可靠的 Type-C 数据线连接电源或电脑，拨动电源开关。",
        "等待启动画面和传感器初始化结束。若显示传感器错误，先重启并检查供电。",
        "对准温差明显、表面哑光的目标，保持目标稳定并覆盖多个热像素。",
        "长按按钮 A 选择容易辨认的伪彩；需要稳定画面时把二次滤波设为轻度。",
        "短按按钮 B 拍照；长按按钮 B 进入相册检查结果。",
        "需要电脑/Android 投屏时保持 USB 连接，打开 BananaThermal Studio 并执行自动连接。",
    ]:
        b.number(step)


def add_device_photos(b: ManualBuilder) -> None:
    b.h1("3 拍照、相册与测温光标")
    b.h2("3.1 拍照与设备相册")
    b.body("短按按钮 B 会触发拍照并冻结当前画面；长按按钮 B 进入设备内部相册。相册文件保存在 LittleFS 中，容量取决于固件分区和设备 Flash。")
    b.bullet("左/右拨动：上一张 / 下一张。")
    b.bullet("按钮 A：第一次显示“确认删除”，再次按按钮 A 执行删除；按其他键取消。")
    b.bullet("按钮 B：退出相册并返回实时画面。")
    b.bullet("照片预览会应用拍摄时保存的热像缩放和偏移；旧照片缺少这些字段时按兼容方式显示。")
    b.callout("先下载再格式化", "“硬盘管理 → 格式化硬盘”会清空 LittleFS 中的照片和大部分配置。格式化前先到上位机“图库 → 设备图库”下载需要保留的照片。", "risk")

    b.h2("3.2 测温光标")
    b.table(
        ["模式", "怎么操作", "适合场景"],
        [
            ["单点", "选择中心、手动、最热或最冷；手动模式通过触摸固定位置", "持续观察一个目标"],
            ["多点", "在画面上依次点击，最多保留 5 个点；按钮 A 可清空", "比较多个部位或多个器件"],
            ["热点 / 冷点追踪", "在当前可见热像区域内自动寻找最高/最低温并平滑位置", "快速定位异常发热或低温区"],
        ],
        [1500, 4300, 3560],
    )
    b.body("缩放、平移、方屏裁剪和旋转都会改变“当前可见区域”。小目标若只覆盖一个热像素的一部分，温度会被周围背景稀释；屏幕放大不会增加传感器真实分辨率。")


def menu_section(
    b: ManualBuilder,
    title: str,
    intro: str,
    rows: list[list[str]],
    tip: str | None = None,
    warn: str | None = None,
    *,
    page_break: bool = False,
) -> None:
    b.h2(title, page_break=page_break)
    b.body(intro)
    b.table(["参数 / 操作", "范围或选项", "怎么用"], rows, [1900, 2550, 4910], compact=True)
    if tip:
        b.callout("建议", tip, "tip")
    if warn:
        b.callout("注意", warn, "warn")


def add_menus_part1(b: ManualBuilder) -> None:
    b.h1("4 设备详细菜单：显示与测温")
    b.body("在主画面按下菜单拨片进入详细菜单。左右选择栏目，按下进入；在子项上按下切换编辑状态，左右改值；移到栏目末尾“返回”并按下时，固件保存设置并回到上一级。")
    menu_section(
        b,
        "4.1 亮度",
        "只控制 LCD 背光，不改变温度数据、照片内容或 USB 推流。",
        [["屏幕亮度", "0–256，步长 1；默认 255", "0 看起来像关屏但设备仍运行；256 显示 MAX 并保持全亮。"]],
        tip="室内降低亮度可减少刺眼和功耗；阳光下再提高。",
    )
    menu_section(
        b,
        "4.2 卡尔曼滤波",
        "对每个热像素做时间平滑，用上一帧状态抑制温度抖动。它和同一帧内的空间滤波不同。",
        [
            ["Q 过程噪声", "最小 0.01，步长 0.01；默认 0.10", "调大响应更快、平滑更弱。"],
            ["R 测量噪声", "最小 0.01，步长 0.01；默认 1.75", "调大更平稳，但拖尾更明显。"],
            ["每图滤波次数", "1–8；默认 4", "同一新帧重复收敛；调大响应更快但增加计算量。"],
        ],
        tip="静态闪烁先小幅增大 R；移动拖尾先减小 R 或增大 Q。一次只改一项。",
        warn="Q、R 在设备菜单中没有合理上限保护，极端数值会让滤波失去意义。",
    )
    menu_section(
        b,
        "4.3 发射率",
        "发射率描述材料表面发出红外辐射的能力。设备显示 95% 时，计算使用 ε=0.95。",
        [
            ["发射率", "0–100%，步长 1%；默认 100%", "人体皮肤、哑光涂层常从 95–98% 起试；亮面金属易受反射影响。"],
        ],
        tip="低发射率表面可在允许时贴已知高发射率胶带，待温度一致后测胶带。",
        warn="不要把 0% 当正常测温设置；发射率设错会造成系统性温差，滤波无法修正。",
    )
    menu_section(
        b,
        "4.4 温度校准",
        "对全部温度应用线性公式：校准后 = 原温度 × 权重 + 偏移。",
        [
            ["权重", "默认 1.00，步长 0.01", "修正比例误差。"],
            ["偏移", "默认 0.00℃，步长 0.01℃", "修正固定偏差。"],
        ],
        tip="至少使用两个跨度足够大的可靠参考点；推荐用上位机“温度线性校准”向导取样、拟合并原子写入。",
        warn="错误校准会影响本机、照片和上位机收到的所有温度。恢复默认：权重 1.00、偏移 0.00。",
    )
    menu_section(
        b,
        "4.5 硬盘管理",
        "这里的“硬盘”是设备内部 Flash 上的 LittleFS，不是电脑硬盘。",
        [
            ["已用 / 总空间", "只读，界面标为 Kb", "实际数值更接近 KiB，用于判断剩余照片容量。"],
            ["格式化硬盘", "立即格式化 LittleFS", "仅在文件系统异常或明确需要清空时使用。"],
        ],
        warn="格式化会删除照片和大部分配置；当前界面可能在底层失败时仍显示成功，完成后必须重启并复查空间与读写。",
    )


def add_menus_part2(b: ManualBuilder) -> None:
    b.h1("5 设备详细菜单：光标、伪彩与触控")
    menu_section(
        b,
        "5.1 测温光标",
        "配置单点、多点和自动冷热追踪的显示方式。",
        [
            ["测温模式", "单点 / 多点；默认单点", "多点最多 5 个触摸点；切换模式会清空旧多点。"],
            ["最高温 / 最低温跟踪", "开 / 关；默认开 / 关", "在当前可见区域自动寻找热点或冷点。"],
            ["跟踪显示", "仅光标 / 光标+温度", "控制自动冷热追踪旁是否显示数值。"],
            ["光标跟踪", "最冷 / 最热 / 中心 / 手动；默认中心", "决定主光标位置来源。"],
            ["测温 / 高温 / 低温字号", "1× / 1.5× / 2×；默认 1.5×", "按屏幕遮挡与可读性选择。"],
            ["数字背景", "透明 / 黑色 / 白色；默认黑色", "让数字在不同伪彩上更清楚。"],
        ],
        tip="目标应覆盖多个热像素；放大后的屏幕像素不等于新增测温点。",
    )
    menu_section(
        b,
        "5.2 伪彩调整",
        "伪彩只改变颜色映射，不改变浮点温度。设备色板与上位机色板可独立选择。",
        [
            ["当前 / 启动伪彩", "经典、鲜艳、热值、翠绿、地狱、白热、黑热", "当前伪彩立即生效；启动伪彩决定下次开机。"],
            ["温度图例", "开 / 关；默认开", "控制整条色标。"],
            ["图例位置", "右上、右下、左上、左下、左中、右中、自定义、信息栏", "自定义时可拖动；信息栏遮挡最少。"],
            ["图例透明度", "0 / 25 / 50 / 75 / 100%；默认 75%", "100% 只隐藏卡片底色，彩条和数值仍可显示。"],
            ["图例详情", "开 / 关；默认开", "控制最高/最低、单位和中间刻度。"],
        ],
        tip="同一绝对温度在不同自动温区下可能显示不同颜色，比较时一定同时看图例数值。",
        warn="当前版本“启动伪彩”左右切换可能都向同一方向循环；遇到时继续循环到目标色板。",
    )
    menu_section(
        b,
        "5.3 触屏调整",
        "控制触摸输入、手势和坐标校准；关闭触摸后实体按键仍可用。",
        [
            ["手势检测", "开 / 关；默认关", "下滑进菜单、上滑退出、左右滑等同左右键、点击等同确认。"],
            ["触屏检测", "开 / 关；默认开", "关闭后停用触摸测温和图例拖动。"],
            ["触屏校准", "依次点击左上和右下角点", "校准后检查四角和中心；偏差更大时用实体键重做。"],
        ],
        tip="校准时用细而不尖锐的触控工具，设备保持稳定。图例拖动不会同时触发滑动手势。",
    )
    menu_section(
        b,
        "5.4 全屏调整",
        "切换本机 LCD 的 320×240 全屏与 240×240 方屏。两种模式分别保存热像和摄像头对齐参数。",
        [
            ["全屏", "开：320×240；关：240×240", "切换时固件保存旧模式配置、加载新模式配置并重新应用视差。"],
        ],
        tip="先选定常用模式，再分别校准摄像头与热像位置；调好全屏并不会自动调好方屏。",
    )


def add_menus_part3(b: ManualBuilder) -> None:
    b.h1("6 设备详细菜单：摄像头、融合与滤波")
    menu_section(
        b,
        "6.1 摄像头",
        "校准可见光裁剪窗口和采集时序；只有摄像头初始化成功时可用。",
        [
            ["全屏 / 方屏 X、Y", "每次按“偏移步长”移动", "改变当前屏幕模式的可见光裁剪位置。"],
            ["全屏 / 方屏缩放", "OV2640 约 240–600；OV3660 约 320–1536", "数值越大通常视野越广、目标越小。"],
            ["热像 X、Y", "-100–100 屏幕像素", "移动当前模式热像显示位置，做少量像素级对齐。"],
            ["偏移步长", "1–128；默认 4", "控制可见光 X/Y 每次调整量。"],
            ["摄像头时钟", "OV3660：2–12；OV2640：0–7", "排查竖纹或不稳定时调节，可能降低帧率。"],
            ["垂直同步", "开 / 关；默认关", "等待完整帧边界可减少撕裂，但改变吞吐。"],
        ],
        tip="对齐顺序：关闭自动视差 → 选屏幕模式 → 调可见光缩放 → 调可见光 X/Y → 少量调热像 X/Y。",
        warn="OV2640 摄像头时钟在当前实现中不一定完整持久化，重启后应重新确认。",
    )
    menu_section(
        b,
        "6.2 融合",
        "把热像颜色与可见光结构组合；两路画面未对齐时会出现重影。",
        [
            ["融合模式", "关闭 / 混合 / 边缘 / 半采样；默认混合", "关闭只看热像；混合看纹理；边缘保留轮廓；半采样降低算力。"],
            ["融合透明度", "0–255，步长 5；默认 126", "0 更偏可见光，255 更偏热像。"],
            ["边缘阈值", "0–255，步长 5；默认 126", "低值细节多、噪声多；高值只保留强轮廓。"],
            ["边缘颜色", "深灰 / 米白；默认深灰", "按热像背景明暗选择。"],
            ["Gamma", "约 0.2–5.0，步长 0.1；默认 1.0", "小于 1 通常提亮中间调，大于 1 通常压暗。"],
            ["旋转 / 镜像 / 红蓝交换", "0/90/180/270°；水平/垂直；开/关", "适配摄像头安装方向和 RGB565 色序。"],
        ],
        tip="先完成几何对齐，再调透明度、阈值和 Gamma；阈值不能修复位置错位。",
    )
    menu_section(
        b,
        "6.3 滤波设置",
        "控制同一帧内的空间平滑和最小显示温区，不替代卡尔曼时间滤波。",
        [
            ["快速插值", "开 / 关；默认关", "开启后跳过二次空间滤波，速度快但噪声和块状感更明显。"],
            ["二次滤波", "关闭 / 轻度 / 中度 / 重度；默认轻度", "越强越平滑，但小热点与细节更容易被抹平。"],
            ["最小显示温差", "0.1–20.0℃，步长 0.1℃；默认 2.0℃", "场景温差很小时撑开伪彩范围，抑制颜色闪烁；不改点温。"],
        ],
        tip="手持巡检可从“卡尔曼默认 + 轻度空间滤波 + 2℃”开始。设备和上位机不要同时把滤波设得过强。",
    )
    menu_section(
        b,
        "6.4 反色设置",
        "用于修正屏幕控制器、软件颜色和 RGB/BGR 顺序；不会改变温度。",
        [
            ["系统反色", "开 / 关；默认关", "修正整屏像底片或电气极性异常。"],
            ["软件反色", "开 / 关；默认关", "对界面和渲染颜色整体取反。"],
            ["屏幕色序", "RGB / BGR；默认 BGR", "红蓝互换时切换；若只有摄像头错色，改融合里的红蓝交换。"],
        ],
        warn="出厂颜色正常时不要修改。一次只切换一项，并用已知红、绿、蓝、黑、白画面判断。",
    )


def add_menus_part4(b: ManualBuilder) -> None:
    b.h1("7 设备详细菜单：传感器、视差与测距")
    menu_section(
        b,
        "7.1 传感器设置",
        "调整热阵列方向、热像视图和采集节奏，并按检测到的传感器显示专用选项。",
        [
            ["热像旋转", "0° / 180°；默认 180°", "旋转温度矩阵及测温坐标。"],
            ["热像缩放", "1.0–2.0×，步长 0.1×；默认 1.0×", "缩小可见视野并放大插值结果，不增加物理像素。"],
            ["速率模式", "极速 / 舒适；默认极速", "主循环延时约 2 ms / 100 ms，不等于传感器硬件刷新率。"],
            ["光学旋转", "0° / 180°", "仅 OV3660 摄像头显示。"],
            ["MLX 速率", "4 / 8 / 16 / 32 Hz；默认 16 Hz", "32 Hz 响应快但对总线、算力和供电要求更高。"],
            ["恢复出厂 Trim", "Heimann 专用维护操作", "仅在 Trim 被错误修改或维护流程明确要求时执行。"],
        ],
        tip="MLX 在 32 Hz 下读帧失败、闪烁或重启时先退回 16 Hz 或 8 Hz。",
        warn="Heimann 恢复 Trim 期间不要断电；这不是普通的“恢复菜单默认值”。",
    )
    menu_section(
        b,
        "7.2 视差调整（可选）",
        "仅 VL53L0X 初始化成功时显示。自动视差根据实测距离在多个 X 补偿锚点之间插值。",
        [
            ["视差模式", "手动 / 自动；默认手动", "手动保持基础对齐；自动随距离补偿。"],
            ["自动基准", "-1000–1000；默认 0", "整体平移所有锚点曲线，适合统一装配偏差。"],
            ["0.10 m 锚点", "-300–300；默认 250", "标定 10 cm 附近的 X 补偿。"],
            ["0.25 m 锚点", "-300–300；默认 112", "标定 25 cm 附近。"],
            ["0.50 m 锚点", "-300–300；默认 60", "标定 50 cm 附近。"],
            ["1.00 m 锚点", "-300–300；默认 40", "标定 1 m 附近。"],
            ["2.00 m 锚点", "-300–300；默认 0", "标定 2 m 及更远。"],
        ],
        tip="所有距离同方向偏时改自动基准；只有某个距离不准时改对应锚点。先校准测距，再标定视差。",
    )
    menu_section(
        b,
        "7.3 距离传感器（可选）",
        "VL53L0X 为自动视差提供距离输入，不应当作精密测距仪。",
        [
            ["实测距离", "只读，单位 m", "50–2000 mm 通常作为有效距离；更远显示超量程状态。"],
            ["调节步长", "自适应 / 0.001 / 0.01 / 0.1 / 1.0", "控制 XTalk 手工增减。"],
            ["XTalk", "0–64 MCPS", "补偿保护窗、外壳或污渍造成的内部反射。"],
            ["开始校准", "约 50 个样本，通常约 5 秒", "清洁窗口，对准 1.5 m 外均匀墙面并保持不动。"],
            ["清除校准", "把 XTalk 设为 0", "移除当前串扰补偿。"],
        ],
        warn="不同设备的 XTalk 不可照抄。补偿过大可能让有效目标变成无效测量。",
    )


def add_alignment_and_measurement(b: ManualBuilder) -> None:
    b.h1("8 双光对齐与正确测温")
    b.figure(
        DUAL_IMAGE,
        "双光、视差与人体表面温度的原理示意",
        6.15,
        "说明热像与可见光如何融合、近距离视差为何明显，以及热成像测得的是体表温度而非体内温度。",
    )
    b.h2("8.1 推荐对齐流程")
    for step in [
        "关闭自动视差，选择全屏或方屏；准备轮廓清晰、温差明显的平面目标。",
        "在中等距离先调可见光缩放，使可见光视场与热像范围大致一致。",
        "调可见光 X/Y 让主要边缘重合，再用热像 X/Y 做少量屏幕像素微调。",
        "切换另一种屏幕模式，重复缩放与位置校准；两种模式参数相互独立。",
        "有距离传感器时，先完成 XTalk 校准，再依次在 0.10、0.25、0.50、1.00、2.00 m 调视差锚点。",
        "来回移动目标检查插值是否连续；最后再选择融合模式和透明度。",
    ]:
        b.number(step)
    b.callout("为什么近距离更容易错位", "热阵列与摄像头上下或左右相隔几厘米，观察距离越近，两个光心形成的夹角越大，因此画面偏移越明显。", "info")

    b.h2("8.2 怎样获得更可靠的温度")
    b.bullet("先设置正确发射率，避免亮面金属、玻璃和强反射环境直接作为参考。")
    b.bullet("让设备、环境和目标温度稳定；保持距离、角度和目标覆盖比例一致。")
    b.bullet("目标尽量覆盖多个热像素；小热点读数会被周围背景平均。")
    b.bullet("需要可追溯精度时，用黑体或可靠接触式标准进行两点以上校准，并用第三点验证。")
    b.callout("人体测温", "热成像测得的是皮肤表面温度，受环境、气流、运动和局部血流影响，不能直接等同体温。只能用于温度分布观察或辅助筛查，异常时应使用合规接触式体温计复测并按医疗建议处理。", "risk")


def add_studio_install_connect(b: ManualBuilder) -> None:
    b.h1("9 BananaThermal Studio 安装与连接")
    b.body("Next BananaThermal Studio 是 Windows 与 Android 共用的上位机，可同时接收热流和可见光流，完成融合、测温标记、趋势记录、设备图库下载与本地录像回放。")
    b.h2("9.1 安装")
    b.table(
        ["平台", "推荐包", "安装说明"],
        [
            ["Windows x64", "BananaThermal-windows-x64-setup.exe", "运行安装程序；若使用便携 ZIP，完整解压后运行 banana_thermal.exe。"],
            ["Android", "BananaThermal-android.apk", "允许安装未知来源应用；连接设备时授权 USB Host。"],
        ],
        [1700, 3500, 4160],
    )
    b.callout("下载稳定版", "日常使用选择 Releases 中带版本号的正式包；CI Artifacts 属于测试版，功能更新但可能不稳定。", "tip")
    b.h2("9.2 连接设备")
    for step in [
        "用支持数据传输的 Type-C 线把热成像设备连接到电脑或 Android 手机，打开设备电源。",
        "启动 Studio。Windows 会自动扫描一次；Android 请点“自动”并在系统弹窗中允许 USB 访问。",
        "波特率可选 115200、230400、460800、921600、1000000。默认先选 115200；只有设备固件明确配置为高速串口时才提高。",
        "自动连接失败时点“刷新”，手动选择 COM/USB 端口，再点“连接”。",
        "看到绿色“已连接”、设备序列号和激活状态后，进入“实时”。",
    ]:
        b.number(step)
    b.table(
        ["控件", "作用", "排查要点"],
        [
            ["端口下拉", "选择 COM 口或 Android USB 设备", "端口不存在：换数据线、USB 口或重新供电。"],
            ["刷新", "重新枚举端口", "更换线缆后先刷新。"],
            ["自动", "扫描并探测兼容设备", "Android 会触发 USB 权限；手动连接后不会被自动扫描抢占。"],
            ["连接 / 断开", "建立或关闭串口会话", "断开前先停止录像或设备图库下载。"],
            ["设备信息", "显示序列号、激活、激活时间和保修截止", "Android 默认折叠，点“设备信息”展开。"],
        ],
        [1700, 3300, 4360],
        compact=True,
    )
    b.callout("未激活设备", "软件会显示激活遮罩。复制序列号，按界面提示向售后获取激活码，粘贴后点“激活”。激活失败时先确认设备保持连接、序列号完整且激活码没有空格。", "warn")
    b.callout("Android 返回键", "手机窄屏中，返回键会优先关闭弹窗、详情页、侧栏或全屏，再返回上一页；平板/宽屏通常保留双栏布局。若系统弹出 USB 权限请求，必须先授权才能自动连接。", "info")


def add_studio_realtime(b: ManualBuilder) -> None:
    b.h1("10 Studio：实时画面与温度记录")
    b.figure(
        STUDIO_REALTIME,
        "Windows 实时双光主页",
        6.3,
        "Studio 实时页包含连接栏、热像与可见光推流开关、主画面、最高最低平均温度、趋势图、渲染参数、拍摄栏和串口控制台。",
    )
    b.h2("10.1 主画面常用操作")
    b.table(
        ["操作", "怎么做", "结果"],
        [
            ["开启热成像", "点“热成像推流”开关", "开始接收温度矩阵并显示 MAX/MIN/AVG。"],
            ["开启可见光", "点“可见光推流”开关", "单独显示或参与融合；关闭热流时可只看可见光。"],
            ["切换测温方式", "点单点/多点图标，再点画面", "单点固定一个位置；多点可放多个温度标签。"],
            ["清理光标", "点清理图标", "移除当前单点和多点标记。"],
            ["全屏", "点四角全屏图标", "Windows 进入软件内全屏视图；Android 手机转横屏并进入沉浸模式。"],
            ["拍照", "点相机图标", "把当前启用的热/可见光槽位保存为 .btpkg 照片。"],
            ["录像", "点摄像机图标开始，再点一次停止", "把每帧槽位写入 .btpkg 视频；录制不必中断实时显示。"],
        ],
        [1700, 3300, 4360],
        compact=True,
    )
    b.callout("空帧确认", "若热成像和可见光两路都关闭，拍照或录像只会得到时间戳空帧；软件会弹窗确认。通常应至少打开一路推流。", "warn")
    b.callout("录像槽位是逐帧记录的", "录制期间可以单独开关热像或可见光。每一帧都会记录当时实际存在的槽位，因此同一个视频可依次包含双路、纯热、纯可见或空帧；回放时不要把中途缺一路误判为文件损坏。", "info")

    b.h2("10.2 调整画面")
    b.table(
        ["面板", "可调项目", "建议"],
        [
            ["颜色映射", "jet、hot、inferno、magma、plasma、viridis、cividis、turbo、gray、cool、rainbow；线性 / S 曲线", "共 11 套，只影响 Studio 显示，不修改设备温度。"],
            ["上采样", "倍率；最近、双线性、双三次", "双三次更平滑但更耗算力；原始热像分辨率不变。"],
            ["滤波", "双边滤波、温度卡尔曼、像素卡尔曼、σ 空间、σ 亮度", "两端滤波不要都设很强，以免延迟和小目标丢失。"],
            ["融合", "关闭、混合、边缘叠加；强度/阈值/粗细/Gamma", "先确认可见光与热像对齐，再调视觉效果。"],
            ["显示", "最高/最低追踪、温度图例方向/停靠/透明度/数字刻度", "这些是上位机本地显示设置，不会写回设备菜单。"],
        ],
        [1500, 4400, 3460],
        compact=True,
    )

    b.h2("10.2.1 可见光悬浮窗")
    b.body("双路推流时，可把可见光作为悬浮小窗叠放在热像上。用可见光窗口开关显示或隐藏；按住标题/拖动区域移动位置，拖动边缘或角点改变尺寸。关闭小窗只影响当前显示，不会自动关闭可见光推流，也不会删除正在录制的可见光槽位。")

    b.h2("10.3 温度趋势与连续记录")
    b.body("收到热像帧后，软件按所选频率把最高、最低、平均、固定单点、各多点及多点平均写入当前会话。普通视图和全屏视图共用同一记录。")
    b.bullet("暂停 / 恢复：只暂停趋势记录，热成像画面继续刷新。")
    b.bullet("记录频率：5 次/秒、2 次/秒、1 次/秒、每 2/5/10 秒一次；常规监测推荐 1 次/秒。")
    b.bullet("保存并新建：把当前会话归档到“历史”，下一帧开始新会话。")
    b.bullet("导出：可多选 CSV 原始数据、PNG 趋势图、JSON 结构化数据。")
    b.callout("导出差异", "CSV 适合 Excel/Python 分析；PNG 适合汇报；JSON 保留时间戳、设备序列号、坐标和每个测温点，适合系统集成。", "info")

    b.h2("10.4 串口控制台")
    b.body("展开“串口控制台”可查看设备日志，也可输入原始命令并回车或点“发送”。常见只读诊断命令包括 help、GetSysInfo、thermal view、calibration get、filter show 和 badpix list。普通用户不应发送来源不明的写入命令；完整的分级命令表见附录 B。")


def add_studio_galleries(b: ManualBuilder) -> None:
    b.h1("11 Studio：设备图库与软件图库")
    b.h2("11.1 设备图库")
    b.figure(
        STUDIO_DEVICE_GALLERY,
        "设备图库：下载、解析、测温标记与导出",
        6.25,
        "设备图库左侧为设备照片列表，右侧为详情预览，可调整颜色、融合、温度标记并导出 PNG。",
    )
    b.number("连接设备后进入“图库 → 设备图库”，点“刷新”获取设备内部照片列表。进入图库时软件会暂停实时流，离开时恢复原推流状态。")
    b.number("点照片开始下载。首次完成后会保存原始文件和缓存；同一会话再次打开可秒开。")
    b.number("在详情中调整颜色映射、映射曲线、融合模式、混合度、Gamma 或边缘参数。")
    b.number("点画面添加温度标记；可切换可见光小窗、温度叠加或清空全部标记。")
    b.number("点“导出 PNG”。导出图会烘焙当前颜色、融合、MAX/MIN/AVG 和温度标记。")
    b.bullet("可见光窗口：显示/隐藏可见光小窗；温度叠加：显示/隐藏 MAX/MIN/AVG 与点温标记。")
    b.bullet("元数据：查看照片格式、热阵列尺寸、可见光尺寸、拍摄参数及可用槽位；旧格式缺少的字段会标为不可用。")
    b.bullet("下载队列：同一会话优先复用内存缓存，跨会话可复用磁盘缓存；刷新列表只清空会话状态。")
    b.bullet("JPEG 设备照片会按原始图像保存，并可作为 JPG 使用；Android 导出时可选择写入系统相册。")
    b.callout("设备端删除", "Studio 的设备图库主要负责下载和导出；删除设备内部照片请在设备相册中二次确认执行。刷新设备列表会清空单次会话缓存，但不会删除磁盘缓存。", "info")

    b.h2("11.2 软件图库")
    b.figure(
        STUDIO_SOFTWARE_GALLERY,
        "软件图库：.btpkg 照片/视频回放与导出",
        5.7,
        "软件图库可管理本地 btpkg 数据包、逐帧播放视频、调整渲染参数、显示温度追踪并导出 PNG 或 MP4。",
    )
    b.table(
        ["操作", "说明"],
        [
            ["刷新 / 选择", "重扫本地 SoftwareGallery 目录；点条目打开详情。"],
            ["重命名 / 备注", "通过更多菜单修改文件名或添加场景、标签等备注。"],
            ["多选 / 删除", "进入多选后全选或批量删除；删除 .btpkg 不可恢复。"],
            ["照片导出", "把当前帧按现有渲染参数导出 PNG。"],
            ["视频播放", "播放/暂停，拖动进度条跳帧；到末帧后再次播放会从头开始。"],
            ["视频温度追踪", "可勾选热点、冷点；效果会写入导出视频。"],
            ["温度标签", "点画面添加标签；视频中位置固定、温度随每帧重新采样。"],
            ["导出 MP4", "Windows 使用随包 ffmpeg；Android 使用系统硬件编码。找不到 ffmpeg 时回退逐帧 PNG。"],
        ],
        [2300, 7060],
        compact=True,
    )
    b.bullet("每帧可为纯热、纯可见、双路或空槽位；时间轴按每帧 slot mask 显示实际内容。")
    b.bullet("详情页可显示/隐藏可见光、温度叠加和标签；这些开关只改变回放与导出效果，不改写原始 .btpkg。")
    b.bullet("照片与视频的“导出当前帧”只生成当前画面；整段视频使用 MP4，缺少编码器时批量导出 PNG 帧。Android 也可保存到系统相册。")
    b.callout("文件位置", "设置页“图库下载路径”决定设备图库 raw/exports 目录；实时拍摄和录像进入软件图库目录。恢复出厂设置不会删除已下载文件或 .btpkg 数据包。", "tip")


def add_studio_history_settings(b: ManualBuilder) -> None:
    b.h1("12 Studio：历史、校准、设置与更新")
    b.h2("12.1 温度历史")
    b.table(
        ["区域", "可执行操作"],
        [
            ["会话列表", "搜索名称或设备序列号；按全部、今天、近 7 天、近 30 天筛选。"],
            ["批量管理", "全选当前筛选结果、清空选择、永久删除所选会话。"],
            ["单个会话", "刷新、重命名、结束并归档、删除、导出 CSV/PNG/JSON。"],
            ["趋势分析", "查看全部、近 1 分钟、近 10 分钟、近 1 小时；选择最高、最低、平均、固定单点、多点均值或各测温点曲线。"],
            ["导入", "选择此前导出的温度历史 JSON，恢复为本地历史会话。"],
        ],
        [2100, 7260],
    )
    b.callout("图表抽样", "超长会话在界面中最多等距绘制约 1200 点以保证流畅；CSV/JSON 导出仍保留全部原始记录。", "info")

    b.h2("12.2 温度线性校准向导")
    for step in [
        "连接设备并打开热成像推流，在设置页点“打开校准向导”。",
        "先读取设备当前增益 k、偏移 b 和存储状态。",
        "让均匀恒温目标覆盖中央取样框，待稳定后输入参考温度，点“采集中央区域”。",
        "至少添加两个不同温点；推荐跨度不小于 10℃，并避免单点标准差超过约 0.2℃。",
        "检查拟合公式、残差和安全范围，点“写入并保存到设备”。",
        "用第三个未参与拟合的温点复核。需要撤销时点“恢复默认系数”。",
    ]:
        b.number(step)
    b.callout("写入范围", "新协议只接受增益 0.5–1.5、偏移 -100–100℃，并原子保存。旧固件可能分两次写入，建议先升级。原有照片中的历史温度不会被改写。", "warn")
    b.callout("拟合质量怎么看", "优先看温点跨度、RMSE（均方根误差）和每个点的残差。跨度太小会放大斜率误差；RMSE 或单点残差偏大时，重新稳定目标并点刷新/重新采集，不要勉强写入。", "tip")

    b.h2("12.3 设置页")
    b.table(
        ["设置", "作用"],
        [
            ["界面缩放 DPI", "调整文字和控件密度，范围 0.8–1.6。"],
            ["主题模式", "跟随系统、白天或夜间。"],
            ["响应式断点", "桌面主区宽度超过阈值时切换宽屏布局，默认 1100。"],
            ["窗口尺寸", "桌面应用预设/自定义宽高、最大化或恢复；最小 600×400。"],
            ["图库下载路径", "设置设备图库原始文件与 PNG 导出的根目录。"],
            ["软件更新", "手动检查；启用后每次启动检查，失败后每分钟重试。"],
            ["恢复出厂设置", "重置主题、缩放、断点、控制台、下载路径、更新和窗口尺寸；不删除下载文件。"],
            ["关于", "查看软件用途与版本信息。"],
        ],
        [2300, 7060],
    )
    b.h2("12.4 软件更新")
    b.body("更新系统优先使用 GitHub 官方发布；受限网络下可切换多个镜像，并校验文件大小与 SHA-256。发现新版本后可查看更新内容、忽略版本、稍后提醒或下载并安装。")
    b.bullet("Android：下载 APK、校验后交给系统安装器。新旧 APK 签名不一致时可能需要先备份并卸载旧版。")
    b.bullet("Windows 安装版：下载并启动 Setup 覆盖升级；便携 ZIP：退出应用后解压覆盖旧目录。")
    b.bullet("镜像只有一个有效响应时，软件会明确提示“未交叉验证”；如有疑虑应稍后重试或前往官方发布页核对。")
    b.bullet("下载中可以取消；失败后可重试或切换镜像。Windows 启动安装器后若应用仍在运行，应按提示退出；安装器窗口未出现时可从下载完成页重新打开。")


def add_firmware_upgrade(b: ManualBuilder) -> None:
    b.h1("13 固件下载与升级")
    b.callout("升级前备份", "先下载设备照片，并记录发射率、温度校准、触摸校准、摄像头/热像对齐和视差锚点。Flash 分区变化可能让旧 LittleFS 无法读取。", "risk")
    b.h2("13.1 选择正确 UF2")
    b.table(
        ["文件后缀", "硬件", "何时选择"],
        [
            ["sketch3m_fs5m.uf2", "8 MB Flash，300 MHz", "默认正常版，性能优先。"],
            ["sketch3m_fs5m_stable.uf2", "8 MB Flash，220 MHz", "可见光竖纹、随机重启或高频不稳定时优先测试。"],
            ["sketch1m_fs1m.uf2", "2 MB Flash", "旧款或小容量设备；照片和配置空间更小。"],
        ],
        [3300, 2700, 3360],
    )
    b.callout("不要猜 Flash 容量", "不能根据 UF2 文件大小推断硬件容量。刷入错误分区可能导致无法启动、文件系统不匹配或原数据不可用。", "warn")
    b.h2("13.2 刷写步骤")
    for step in [
        "从正式 Release 下载与设备匹配的 UF2，同时下载 SHA256SUMS、SHA256SUMS.sig 和允许签名者文件；完成校验。",
        "关闭设备并断开 USB。按住按钮 A（BOOTSEL），再按复位键；先松开复位键，看到 RP2040 可移动磁盘后再松开按钮 A。",
        "把 UF2 复制到 RP2040 可移动磁盘，等待磁盘自动断开和设备重启。复制过程中不要拔线。",
        "首次启动给文件系统初始化或配置迁移留出时间，不要立刻断电。",
        "检查版本、传感器识别、全屏/方屏、热像、可见光、融合和照片。",
        "用最新 Studio 连接并检查 thermal view、实时投屏及校准参数。",
    ]:
        b.number(step)
    b.callout("不同外壳", "若实物没有明确标注 BOOT/RESET，按设备批次资料操作，不要拆机短接未知焊盘。", "warn")


def add_troubleshooting(b: ManualBuilder) -> None:
    b.h1("14 故障排查")
    b.table(
        ["现象", "优先检查", "下一步"],
        [
            ["刷写后不启动", "UF2 是否匹配 8 MB/2 MB；供电和线缆；是否使用稳定版", "重新进入 Bootloader，刷回已校验的正确版本。"],
            ["电脑找不到设备", "是否为数据线；端口是否被占用；设备是否开机", "换 USB 口/线，Studio 点刷新或自动，必要时重启。"],
            ["连接后没有热像", "热成像推流开关；串口波特率 115200；传感器初始化", "展开控制台查看日志，断开重连。"],
            ["投屏一会儿停止", "主线程是否被阻塞；USB 稳定性；是否在切换横竖屏", "退出全屏后重连；更新 Studio，减少同时下载/录像负载。"],
            ["方屏被拉伸/裁切", "Studio 是否收到 thermal view；screen_mode、thermal_height、scale", "更新 Studio；不要把全屏缩放和偏移套到方屏。"],
            ["Heimann 方屏无法解析", "thermal_height 应为 32、thermal_count 应为 1024", "使用支持 Thermal View v2 的最新 Studio。"],
            ["可见光竖纹或随机重启", "USB 供电；摄像头时钟/垂直同步；固件频率", "测试 8 MB stable 版，分别测试纯热、纯可见、融合。"],
            ["融合重影", "两路旋转/镜像；缩放；X/Y；视差模式和距离", "按“中距离基础对齐 → 测距 → 锚点”顺序重做。"],
            ["画面闪烁", "卡尔曼 Q/R、空间滤波、最小显示温差", "静态闪烁增大 R 或显示温差；移动拖尾反向调整。"],
            ["温度明显不准", "发射率、反射、距离、目标覆盖、校准参数", "恢复默认校准后用可靠参考源重新标定。"],
            ["设备图库下载失败", "是否仍在推流；连接是否稳定；照片文件是否损坏", "刷新列表、重试下载；换线后重新连接。"],
            ["MP4 导出失败", "Windows 随包 ffmpeg 是否完整；磁盘空间；文件权限", "使用安装版/完整便携包；软件会在缺少 ffmpeg 时回退逐帧 PNG。"],
        ],
        [2200, 3750, 3410],
        compact=True,
    )
    b.h2("14.1 提交问题时准备")
    b.bullet("设备型号、Flash 容量、固件完整版本和 UF2 变体。")
    b.bullet("热阵列型号、摄像头型号、是否带 VL53L0X、当前全屏/方屏模式。")
    b.bullet("Studio 版本、操作系统、波特率、thermal view 响应和最短复现步骤。")
    b.bullet("已删除序列号、激活信息、密钥和其他敏感字段的日志或截图。")
    b.callout("不要公开敏感信息", "序列号、激活码、私钥、签名材料和未公开诊断数据不要发布到公开 Issue；安全问题应按仓库 SECURITY.md 私下报告。", "risk")

    b.h1("附录 A 快速速查", page_break=False)
    b.table(
        ["想完成的事", "最短路径"],
        [
            ["换伪彩", "长按按钮 A；或详细菜单 → 伪彩调整。"],
            ["拍照", "短按按钮 B；Studio 中点相机图标会另存为 .btpkg。"],
            ["进入设备相册", "长按按钮 B。"],
            ["清理测温点", "短按按钮 A 分层清理；Studio 点清理图标。"],
            ["快速调亮度/缩放/融合", "按快捷菜单键 → 左右选项 → 按下进入编辑 → 左右调值。"],
            ["两路画面对齐", "摄像头缩放 → 可见光 X/Y → 热像 X/Y → 测距校准 → 视差锚点。"],
            ["连接上位机", "数据线连接 → Studio 自动 → 115200 → 已连接 → 开启热像/可见光推流。"],
            ["导出温度数据", "实时趋势或历史 → 导出 → CSV / PNG / JSON。"],
            ["下载设备照片", "Studio → 图库 → 设备图库 → 刷新 → 选择照片 → 导出 PNG。"],
            ["导出录像", "Studio → 软件图库 → 选择视频 → 导出 MP4。"],
            ["恢复温度校准", "设备菜单权重 1.00/偏移 0.00；或 Studio 校准向导 → 恢复默认。"],
            ["进入 UF2 模式", "关机/断 USB → 按住按钮 A → 按复位 → 先松复位 → 出现磁盘后松 A。"],
        ],
        [2900, 6460],
        compact=True,
    )
    b.h2("附录 A.1 资料与版本依据")
    b.body("本手册内容来自 Banana Thermal Firmware 发布/用户文档、3inch2-Heimann-dual-thermal 固件源码、NextBananaThermalStudio 上位机源码，以及 3.2 寸设备硬件线稿与旧版说明材料。设备行为冲突时，以当前源码和对应 Release Notes 为准。")
    b.callout("维护建议", "后续发布新固件或上位机版本时，优先复核按键状态机、菜单项目数与范围、thermal view/photo 协议版本、Studio 导出格式和截图，再更新本手册版本号。", "tip")


def add_encyclopedia_appendices(b: ManualBuilder) -> None:
    b.h1("附录 B 维护与串口命令", page_break=True)
    b.callout("服务人员功能", "本附录面向排障、校准和开发维护。先用只读命令确认状态，再执行单一写入并立即回读验证。激活算法、私钥、签名材料和寄存器级危险写入不属于公开说明书。", "warn")
    b.h2("附录 B.1 安全排障顺序")
    for step in [
        "停止录像、设备图库下载和固件升级，保存当前工作。",
        "执行 GetSysInfo、thermal view、calibration get、filter show 和 badpix list，并把回显脱敏保存。",
        "一次只执行一个明确的写入命令；不要粘贴来源不明的一长串命令。",
        "立即用对应只读命令回读，再重启设备验证是否持久化。",
        "若行为变差，恢复先前记录值；涉及 clean、激活或未知写入时停止操作并联系维护人员。",
    ]:
        b.number(step)
    b.h2("附录 B.2 命令风险分级")
    b.table(
        ["级别", "命令或命令族", "用途与边界"],
        [
            ["只读 / 推荐", "help · GetSysInfo · ucode · thermal view", "查看帮助、系统信息、识别信息和当前热像视图 JSON。thermal view 是 Studio 解析尺寸与屏幕模式的重要依据。"],
            ["只读 / 推荐", "calibration get · filter show · badpix list", "读取线性校准、滤波状态和坏点表；适合提交问题前记录。"],
            ["只读 / 推荐", "fusion · cam", "不带写入参数时查看融合与摄像头状态；具体回显随固件版本变化。"],
            ["持久化写入", "calibration set <gain> <offset> · calibration reset", "设置或恢复温度线性校准。gain 仅 0.5–1.5，offset 仅 -100–100℃；写后回读。"],
            ["兼容写入", "cali -w0 / -b0 / -show · save", "旧固件分项校准与保存路径；新版本优先使用 calibration 命令或 Studio 向导。"],
            ["滤波写入", "filter set … · kalman set_* …", "修改空间/时间滤波参数。先记录旧值，一次只改一项，出现异常立即恢复。"],
            ["坏点修正", "badpix add/del/clear/on/off/test", "维护坏像素表、启停补偿或测试。add/del/clear 会改变持久配置，必须记录坐标并验证邻域插值。"],
            ["图像写入", "fusion … · cam …", "调整融合、方向或摄像头相关服务参数。普通对齐优先使用设备菜单，避免命令与菜单状态不一致。"],
            ["传输 / 托管", "check · download <filename> · stream · vstream · stop", "照片列表、下载和热像/可见光推流；通常由 Studio 管理，不要在下载或录像中并发手动发送。"],
            ["破坏性", "clean", "清理设备数据/配置。不可恢复；先下载照片并记录全部校准、对齐和菜单参数。"],
            ["受控流程", "activate <code>", "设备激活入口。只使用售后提供的激活码，不公开序列号、激活码或内部实现。"],
        ],
        [1500, 3400, 4460],
        compact=True,
    )

    b.h1("附录 C 数据流、照片与 .btpkg 格式", page_break=True)
    b.h2("附录 C.1 实时串口帧")
    b.table(
        ["数据", "帧结构", "尺寸 / 说明"],
        [
            ["视图描述", "thermal view → JSON v2", "开始推流前读取；包含 screen_mode、thermal_width/height/count、scale 与偏移等解析依据。"],
            ["热像帧", "BEGIN + T_max/T_min/T_avg + 温度数组 + END", "温度与统计值均为 little-endian float32；MLX 为 32×24=768 点，Heimann 方屏为 32×32=1024 点。"],
            ["可见光帧", "VBEG + width/height/length + RGB565 + VEND", "三个长度字段为 little-endian uint32；全屏常见 120×160，方屏常见 120×120。"],
            ["保活", "约 500 ms 心跳，约 1 s 超时", "停止收到有效帧时 Studio 会判定流中断并进入重连/提示流程。"],
        ],
        [1700, 3650, 4010],
        compact=True,
    )
    b.callout("不要按 LCD 分辨率解析", "320×240 或 240×240 是显示结果，不是温度点数。解析器必须以 thermal view 和实际帧长度为准。", "warn")

    b.h2("附录 C.2 设备照片下载与兼容格式")
    b.table(
        ["格式", "识别 / 内容", "Studio 行为"],
        [
            ["v1 Simple", "Tmax、Tmin + 768 个 float32；共 3080 字节", "按 32×24 纯热照片解析。"],
            ["v1 Full", "Tmax、Tmin、full_screen 标志 + 768/1024 个 float32", "依据数组长度与屏幕标志兼容 MLX/Heimann。"],
            ["HTPH v2/v3", "HTPH 魔数、版本/flags、热像及可选 RGB565", "保留双路与视图元数据；旧字段缺失时使用兼容默认值。"],
            ["JPEG", "文件头 FF D8 FF", "作为原始 JPEG/JPG 下载；没有温度矩阵时不能添加真实点温。"],
            ["传输命令", "check；download <filename>；END FILE DATA", "Studio 负责列表、十六进制数据还原、完整性判断、队列与缓存。"],
        ],
        [1700, 3800, 3860],
        compact=True,
    )

    b.h2("附录 C.3 Studio .btpkg v2")
    b.body(".btpkg 是 Studio 的本地照片/视频数据包，不等同于设备照片格式。文件以 BTPK 开头，版本号为 2，标记 photo 或 video；元数据为定长填充的 JSON。随后按帧存储 slotMask、时间戳 tsMs、可选热像 float32 数组和可见光 PNG，结尾为 BTPK_END。")
    b.bullet("slotMask 按帧记录热像/可见光是否存在，因此同一段录像允许槽位动态变化。")
    b.bullet("照片通常只有一帧；视频可逐帧回放、重渲染和重新导出，原始温度数据不会因色板改变。")
    b.bullet("当前 v2 与早期实验性 v1 不承诺直接兼容；迁移重要数据前先保留原文件，并用对应版本 Studio 验证。")
    b.bullet("PNG/JPG/MP4 是面向查看或分享的导出结果；CSV/JSON 是趋势数据导出。只有保留温度矩阵的格式支持后续重新测温。")

    b.h1("附录 D 功能覆盖索引", page_break=True)
    b.table(
        ["功能域", "本手册位置", "覆盖内容"],
        [
            ["硬件与按键", "第 1–3 章", "部件、状态栏、A/B/拨片/快捷键、触摸手势、拍照、相册、测温光标。"],
            ["设备菜单", "第 4–7 章", "16 项详细菜单、范围、默认行为、持久化、硬件条件与风险提示。"],
            ["双光对齐与测温", "第 8 章", "全屏/方屏、摄像头与热像 X/Y/缩放、测距、锚点、发射率与校准。"],
            ["Studio 连接", "第 9 章", "Windows/Android、USB 权限、五档波特率、激活、响应式界面与返回键。"],
            ["Studio 实时", "第 10 章", "双路推流、11 色板、悬浮窗、滤波、融合、点温、趋势、拍照录像与控制台。"],
            ["图库与文件", "第 11 章、附录 C", "设备图库、队列/缓存/元数据、软件图库、动态槽位、PNG/JPG/MP4 与 .btpkg。"],
            ["历史、校准、更新", "第 12 章", "会话管理、CSV/PNG/JSON、RMSE/残差、设置、镜像校验、取消/重试/安装。"],
            ["发布、排障、维护", "第 13–14 章、附录 B", "UF2 变体与 300/220 MHz、校验、故障矩阵、坏点修正、命令风险分级。"],
        ],
        [1900, 2400, 5060],
        compact=True,
    )
    b.callout("百科全书的边界", "本手册覆盖当前公开固件 v1.7.0 与 Studio v0.4.2 可见功能、维护接口和公开格式。硬件生产测试、私钥/激活算法、未发布实验开关与芯片寄存器级调试不公开；它们不属于终端用户可安全操作的功能。", "info")


def build() -> Path:
    doc = Document()
    bullet_num_id, decimal_num_id = configure_document(doc)
    doc.core_properties.title = "3.2 寸双光融合热成像使用说明书"
    doc.core_properties.subject = "设备操作、菜单配置、上位机使用、固件升级与故障排查"
    doc.core_properties.author = "Banana Thermal"
    doc.core_properties.keywords = "热成像, 双光融合, BananaThermal Studio, RP2040, 使用说明书"
    doc.core_properties.comments = "依据固件 v1.7.0 与 Studio v0.4.2 当前实现整理。"

    add_cover(doc)
    b = ManualBuilder(doc, bullet_num_id, decimal_num_id)
    add_front_matter(b)
    add_device_overview(b)
    add_buttons_and_quick_start(b)
    add_device_photos(b)
    add_menus_part1(b)
    add_menus_part2(b)
    add_menus_part3(b)
    add_menus_part4(b)
    add_alignment_and_measurement(b)
    add_studio_install_connect(b)
    add_studio_realtime(b)
    add_studio_galleries(b)
    add_studio_history_settings(b)
    add_firmware_upgrade(b)
    add_troubleshooting(b)
    add_encyclopedia_appendices(b)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    try:
        print(build())
    except Exception as exc:
        print(f"build failed: {exc}", file=sys.stderr)
        raise
